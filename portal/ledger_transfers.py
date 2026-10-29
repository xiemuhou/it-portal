"""Human-readable flat templates for complete ledger records; staged atomic writes."""
from io import BytesIO,StringIO
import csv
import hashlib
import ipaddress
import json
import zipfile
import sqlalchemy as sa
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .database import TABLES as T,insert_row,row,utcnow,jsonable
from .security import require,fail,audit,digest
from . import ledger as L

VERSION='ledger-9'


def template_version(kind):
    # Wi-Fi has its own revision so unrelated ledger templates remain valid.
    return 'ledger-12' if kind=='vpn' else 'ledger-11' if kind=='wifi' else 'ledger-10' if kind=='circuit' else VERSION


def columns(kind):
    if kind=='wifi':
        fields={f['name']:f for f in L.schema(kind)}
        vlan=fields['vlan_network_ids']
        return [('site_id','区域',fields['site_id']),('vlan_network_ids.*.vlan','VLAN ID',vlan),
            ('vlan_network_ids.*.cidr','网段',vlan),('ssid','SSID',fields['ssid']),
            ('wifi_password','密码',fields['wifi_password']),('auth_method','认证方式',fields['auth_method']),
            ('purpose','用途',fields['purpose']),('coverage','覆盖位置',fields['coverage']),
            ('enabled','启用',fields['enabled']),('notes','备注',fields['notes']),
            ('code','记录编号（系统辅助列）',None),('_group','导入分组（系统辅助列）',None)]
    result=[('code','记录编号（新增留空，更新保留）',None),('_group','导入分组（同一条资料填写相同值）',None)]
    for f in L.schema(kind):
        if f.get('ui_hidden'):continue
        if f['name'] in ('lines','tunnels'):
            for i,prefix in enumerate(('主线','赠送线路') if kind=='circuit' else ('隧道1','隧道2')):
                for child in f['columns']:
                    if child['name']=='line_type':continue
                    result.append((f"{f['name']}.{i}.{child['name']}",prefix+'·'+child['label'],child))
        elif f['type']=='rows':
            for child in f['columns']:
                result.append((f['name']+'.*.'+child['name'],(f['label']+'·' if kind!='vlan' else '')+child['label'],child))
        elif f['type']=='refs' and f['target']=='vlan_network':
            for sub,label in [('scope','网络范围'),('vlan','VLAN ID'),('cidr','网段 CIDR')]:
                result.append((f['name']+'.*.'+sub,f['label']+'·'+label,f))
        else:result.append((f['name'],f['label'],f))
    return result


def human_reference(conn,user,f,key):
    target=f['target']
    if target=='vlan_network':
        key,network_id=key.split(':',1);obj=L.detail(conn,user,key,False)
        n=next((n for n in obj['networks'] if n['network_id']==network_id),None)
        if not n:fail('关联的 VLAN 网段已变化，请刷新后重试',409)
        return {'scope':obj.get('scope_name') or 'default','vlan':obj['vlan_number'],'cidr':n['cidr']}
    if target in L.KINDS:
        obj=L.detail(conn,user,key,False)
        if f['name']=='uplink_device_id' and obj.get('management_ip'):
            peers=L.list_items(conn,user,'device',site_id=obj['site_id'])
            address=device_reference_ip(obj['management_ip'])
            if sum(device_reference_ip(v.get('management_ip'))==address for v in peers)==1:return obj['management_ip']
        return obj['code']
    obj=row(conn,target,key);return obj.get('display_name') if target=='users' else obj['name']


def workbook(conn,user,kind,items=None):
    cs=columns(kind);wb=Workbook();sheet=wb.active;sheet.title='数据';sheet.append([c[1] for c in cs])
    for item in items or []:
        fields={f['name']:f for f in L.schema(kind)}
        count=max([1]+[len(item.get(f['name']) or []) for f in fields.values() if not f.get('ui_hidden') and f['type'] in ('rows','refs') and f['name'] not in ('lines','tunnels')])
        for index in range(count):
            values=[]
            for path,label,f in cs:
                if '.*.' in path:
                    key,_,sub=path.split('.');entries=item.get(key,[]);entry=entries[index] if len(entries)>index else None
                    value=human_reference(conn,user,f,entry)[sub] if entry and f['type']=='refs' else entry.get(sub) if entry else None
                    # Empty collections must remain empty when updating an exported record.
                    if not entries and index==0 and sub==next(p.split('.')[-1] for p,_,_ in cs if p.startswith(key+'.*.')):value='__CLEAR__'
                elif '.' in path:
                    key,pos,sub=path.split('.');entries=item.get(key,[])
                    if kind=='circuit':entries=[next((v for v in entries if v.get('line_type')==t),{}) for t in ('primary','gift')]
                    value=entries[int(pos)].get(sub) if len(entries)>int(pos) else None
                elif path=='_group':value=item.get('code')
                elif f and f['type']=='refs':
                    entries=item.get(path) or [];value=human_reference(conn,user,f,entries[index]) if len(entries)>index else ('__CLEAR__' if index==0 and not entries else None)
                else:value=item.get(path)
                if index>0 and path not in ('code','_group') and '.*.' not in path and not(f and f['type']=='refs' and '.' not in path):value=None
                if f and f['type']=='secret':value=None
                if f and f['type']=='ref' and value:value=human_reference(conn,user,f,value)
                if isinstance(value,bool):value='是' if value else '否'
                if value is not None and not isinstance(value,(str,int,float)):value=str(value)
                values.append(value)
            sheet.append(values)
    for cells in sheet:
        for c in cells:
            if isinstance(c.value,str):c.data_type='s'
    for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='2764E7')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    sheet.row_dimensions[1].height=42
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width=24
        for c in col:c.alignment=Alignment(vertical='top',wrap_text=True)
    if kind=='wifi':
        for position,(path,_,_) in enumerate(cs,1):
            if path in ('code','_group'):sheet.column_dimensions[sheet.cell(1,position).column_letter].hidden=True
    info=wb.create_sheet('模板信息');info.append(['台账类型',kind]);info.append(['版本',template_version(kind)])
    guide=wb.create_sheet('填写说明')
    for value in ['新增编号留空，系统生成；更新时保留导出编号。更新空格保留原值，__CLEAR__ 明确清空选填字段。',
                  '区域填写名称，安装位置直接填写；上联设备可直接填同一区域内唯一的管理 IP，也兼容 型号 · 管理IP、记录编号和名称。同批设备无需按上下联排序。其他关联设备/线路可填编号或唯一名称。',
                  '设备更新模式：记录编号优先；编号留空时按区域 + 管理 IP 匹配已有设备，唯一匹配才更新；没有匹配则新增。仅新增模式会拒绝同区域重复管理 IP。',
                  '每项信息独立一列；多条网段、接入明细或关联资料逐行填写，不需要竖线、@ 或分号拼接。',
                  '同一条资料的多行填写相同记录编号；新增时编号留空，导入分组填相同自定义值（例如 A001）。不同资料使用不同分组。单行资料可不填分组。',
                  '新增 VLAN 也可用相同区域、网络范围和 VLAN ID 自动合并多行。共享字段可只填首行；重复填写时内容必须一致。',
                  '关联 VLAN 的网络范围、VLAN ID、网段 CIDR 分列填写，区域按本地 / 对端区域匹配；网络范围留空默认为 default。',
                  'VPN 隧道关联多条 VLAN 网段时，每个网段填写一行，所有行的“导入分组”必须填写完全相同的值（例如 VPN001）；新增时记录编号保持空白。共享字段只填首行或每行填写相同内容。',
                  '本地或对端 VPN 网段可填写 selected（指定 VLAN 网段）或 any（Any - 0.0.0.0/0）；也可直接填写中文“指定 VLAN 网段”、Any 或 0.0.0.0/0。选择 Any 时该侧 VLAN 网段会被忽略。',
                  '每行一个网段，各自填写 CIDR、网关、网段用途。网段编号、接入编号新增留空，更新保留以维持引用和密码。',
                  'Wi-Fi 模板按区域、VLAN ID、网段、SSID、密码、认证方式、用途、覆盖位置、启用、备注排列；同一 Wi-Fi 关联多个网段时，每行重复填写区域和 SSID。',
                  'Wi-Fi 认证方式填写 不认证 / 密码 / 密码+MAC（也可填 none / password / password_mac）。Wi-Fi、设备、光猫及拨号密码列可导入，导出留空，更新留空保留原密码。',
                  '多个可用 IP、对端网段逐行填写，各字段分列。掩码可填 255.255.255.248 或 29。',
                  '隧道登记状态 up（正常）/ disabled（停用），也可填写中文正常 / 停用；接入连接方式 route（路由）/ bridge（桥接）/ pppoe（拨号）；启用、是否堆叠填 是 / 否。',
                  '线路接入明细每行一个接入，所属线路填 primary / gift（或主线 / 赠送线路）。新增接入编号留空；更新保留。',
                  '多种明细可在同一行各填一条，无需两两组合。更新时提供该资料完整明细；单元格留空保留共享字段，明细首列填 __CLEAR__ 可明确清空整个明细。',
                  '主线、赠送线路及两个隧道使用各自列，不需要额外建表。日期 YYYY-MM-DD。',
                  '影响已有引用的 VLAN 网络修改，请在页面检查关联影响后保存，不通过导入跳过确认。']:
        guide.append([value])
    guide.column_dimensions['A'].width=120
    out=BytesIO();wb.save(out);return out.getvalue()


def device_reference_ip(value):
    try:return str(ipaddress.ip_address(str(value or '').rsplit('·',1)[-1].strip()))
    except ValueError:return None


def resolve(conn,user,kind,data):
    fields={f['name']:f for f in L.schema(kind)}
    for name,f in fields.items():
        if name not in data or not data[name] or f['type'] not in ('ref','refs'):continue
        target=f['target']
        def match(value):
            if target=='vlan_network':
                options=L.network_options(L.list_items(conn,user,'vlan',site_id=data.get(f.get('site_field','site_id'),'')))
                choices=[v['id'] for v in options if value in (v['id'],(v.get('scope_name') or 'default')+'@'+str(v['vlan_number'])+'@'+v['cidr'])]
            elif target in L.KINDS:
                options=L.list_items(conn,user,target,site_id=data.get(f.get('site_field','site_id'),'') if not f.get('global_ref') else '')
                address=device_reference_ip(value) if name=='uplink_device_id' else None
                if address:
                    choices=[v['id'] for v in options if device_reference_ip(v.get('management_ip'))==address]
                    if len(choices)!=1:fail('上联设备管理 IP '+address+(' 在当前区域重复，请使用唯一记录编号' if choices else ' 在当前区域未找到，请检查区域、IP 或同批设备数据'))
                else:choices=[v['id'] for v in options if value.strip() in (v['id'],v['code'],v['name'],((v.get('scope_name') or 'default')+'@'+str(v.get('vlan_number'))) if target=='vlan' else '')]
            else:
                options=L.rows(conn,target)
                choices=[v['id'] for v in options if value in (v['id'],v.get('name'),v.get('display_name'),v.get('login_name')) and (target!='locations' or v['site_id']==data.get('site_id'))]
            if len(choices)!=1:fail(f['label']+'未匹配到唯一资料：'+value)
            return choices[0]
        data[name]=[match(v) for v in (data[name] if isinstance(data[name],list) else L.strings(data[name]))] if f['type']=='refs' else match(str(data[name]))
    return data


def order_device_groups(conn,user,groups):
    """Topological ordering of in-file uplinks, preserving original error row numbers."""
    providers={};dependencies={};ordered=[]
    for index,g in enumerate(groups):
        data=g['data']
        try:
            if data.get('site_id'):data.update(resolve(conn,user,'device',{'site_id':data['site_id']}))
            elif data.get('code'):
                existing=L.rows(conn,'resources',resource_code=data['code'])
                if existing:data['site_id']=existing[0]['site_id']
            address=device_reference_ip(data.get('management_ip'))
            aliases=[('code',str(data['code']).strip())] if data.get('code') else []
            if address:aliases.append(('ip',data.get('site_id'),address))
            for alias in aliases:providers.setdefault(alias,[]).append(index)
        except Exception as exc:
            if not hasattr(exc,'detail'):raise
            g['error']=str(exc.detail)
    for index,g in enumerate(groups):
        data=g['data'];value=data.get('uplink_device_id');address=device_reference_ip(value)
        alias=('ip',data.get('site_id'),address) if address else ('code',str(value or '').strip())
        matches=providers.get(alias,[])
        if len(matches)>1:g['error']='上联设备在同一批次内匹配到多条资料，请检查同一区域的管理 IP 或记录编号'
        dependencies[index]=set(matches) if len(matches)==1 else set()
    remaining=set(range(len(groups)));done=set()
    while remaining:
        ready=[i for i in sorted(remaining) if not dependencies[i]-done]
        if not ready:
            for i in sorted(remaining):
                groups[i]['error']='同批上联设备存在循环关联或关联到自身，请检查管理 IP'
                ordered.append(groups[i])
            break
        for i in ready:ordered.append(groups[i]);remaining.remove(i);done.add(i)
    return ordered


def read_groups(sheet,kind,cs):
    """Collapse worksheet detail rows into one validated proposal per resource."""
    it=iter(sheet.iter_rows())
    if [c.value for c in next(it,[])]!=[v[1] for v in cs]:fail('模板表头不一致')
    groups={};fields={f['name']:f for f in L.schema(kind)}
    for line,cells in enumerate(it,2):
        if line>5001:fail('每批最多 5000 行')
        if not any(c.value is not None for c in cells):continue
        raw={p:c.value for (p,_,_),c in zip(cs,cells) if c.value is not None}
        code=str(raw.get('code') or '').strip();group=str(raw.get('_group') or '').strip()
        identity=('code',code) if code else ('new',group) if group else (
            ('vlan',str(raw.get('site_id') or ''),str(raw.get('scope_name') or 'default'),str(raw['vlan_number']))
            if kind=='vlan' and raw.get('site_id') and raw.get('vlan_number') is not None else
            ('wifi',str(raw.get('site_id') or ''),str(raw.get('ssid') or ''))
            if kind=='wifi' and raw.get('site_id') and raw.get('ssid') else ('line',line))
        g=groups.setdefault(identity,{'line':line,'data':{},'error':None,'fixed':{},'arrays':{}})
        try:
            if any(c.data_type=='f' for c in cells):fail('台账导入不允许公式')
            detail={}
            for path,label,f in cs:
                if path not in raw or path=='_group':continue
                value=raw[path]
                if hasattr(value,'isoformat'):value=value.isoformat()[:10]
                if value=='':continue
                if '.*.' in path:
                    key,_,sub=path.split('.');detail.setdefault(key,{})[sub]=value;continue
                if '.' in path:
                    old=g['fixed'].get(path)
                    if old is not None and str(old)!=str(value):fail(label+'在同一分组中不一致')
                    g['fixed'][path]=value;continue
                if f and f['type']=='refs':
                    entries=g['arrays'].setdefault(path,[])
                    v=[] if value=='__CLEAR__' else [str(value)]
                    if v and v[0] not in entries:entries.extend(v)
                    continue
                if value=='__CLEAR__':value=None
                elif f and f['type']=='bool':
                    if value not in (True,False,1,0,'是','否','true','false'):fail('启用状态须为是或否')
                    value=value in (True,1,'是','true')
                elif f and f['type']=='int':
                    if str(value).strip()!=str(int(value)):fail(label+'须为整数')
                    value=int(value)
                else:value=str(value)
                if path in g['data'] and g['data'][path]!=value:fail(label+'在同一分组中不一致')
                g['data'][path]=value
            for key,entry in detail.items():
                entries=g['arrays'].setdefault(key,[]);f=fields[key]
                first=next(p.split('.')[-1] for p,_,_ in cs if p.startswith(key+'.*.'))
                if entry.get(first)=='__CLEAR__':
                    if len(entry)!=1 or entries:fail(f['label']+'清空标记不能与明细混用')
                    g.setdefault('cleared',set()).add(key);continue
                if key in g.get('cleared',set()):fail(f['label']+'清空标记不能与明细混用')
                if f['type']=='refs':
                    if '__CLEAR__' in entry.values() or not entry.get('vlan') or not entry.get('cidr'):fail(f['label']+'须填写 VLAN ID 和网段 CIDR')
                    value=str(entry.get('scope') or 'default')+'@'+str(entry['vlan'])+'@'+str(entry['cidr'])
                    if value not in entries:entries.append(value)
                else:
                    entry={k:None if v=='__CLEAR__' else str(v) for k,v in entry.items()}
                    if key=='connections' and entry.get('line_type') in ('主线','赠送线路'):entry['line_type']={'主线':'primary','赠送线路':'gift'}[entry['line_type']]
                    entries.append(entry)
        except Exception as exc:
            if hasattr(exc,'detail'):message=str(exc.detail)
            elif isinstance(exc,(ValueError,TypeError)):message='字段格式无效'
            else:raise
            g['error']='第 '+str(line)+' 行：'+message
    for g in groups.values():
        data=g['data'];data.update(g['arrays'])
        for path,value in g['fixed'].items():
            name,index,child=path.split('.');entry=data.setdefault(name,[{},{}])[int(index)]
            entry[child]=None if value=='__CLEAR__' else str(value)
            if kind=='circuit':entry['line_type']=('primary','gift')[int(index)]
        for name in ('lines','tunnels'):
            if name in data:data[name]=[e for e in data[name] if any(v not in (None,'') for k,v in e.items() if k!='line_type')]
    return list(groups.values())


def preview(conn,user,kind,raw,filename,mode):
    require(conn,user,'resource.import');L.schema(kind)
    if mode not in ('create_only','upsert'):fail('导入模式无效')
    if not filename.lower().endswith('.xlsx') or len(raw)>20*1024*1024:fail('仅支持 20 MB 以内 XLSX')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if sum(v.file_size for v in z.infolist())>100*1024*1024 or any('vbaproject' in v.filename.lower() for v in z.infolist()):fail('文件过大或包含宏')
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except (ValueError,KeyError,zipfile.BadZipFile):fail('无法读取台账模板')
    staged=[];seen=set();cs=columns(kind);device_tx=None
    try:
        if set(wb.sheetnames)!={'数据','模板信息','填写说明'} or dict(wb['模板信息'].values)!={'台账类型':kind,'版本':template_version(kind)}:fail('请使用本台账最新标准模板')
        groups=read_groups(wb['数据'],kind,cs)
        if kind=='device':
            groups=order_device_groups(conn,user,groups)
            device_tx=conn.begin_nested()
        for position,group in enumerate(groups):
            line=group['line'];data=group['data'];error=None;target=None;expected=None
            try:
                if group['error']:fail(group['error'])
                if data.get('code'):
                    matches=conn.execute(sa.select(T['resources']).where(T['resources'].c.resource_code==data['code'])).mappings().all()
                    if matches:
                        if mode=='create_only':fail('编号已存在，请选择更新模式')
                        if matches[0]['resource_kind']!=kind:fail('编号对应的台账类型不一致')
                        target=matches[0]['id'];expected=matches[0]['version']
                if kind=='device' and not target and data.get('management_ip'):
                    address=device_reference_ip(data['management_ip'])
                    matches=[v for v in L.list_items(conn,user,'device',site_id=data.get('site_id','')) if address and device_reference_ip(v.get('management_ip'))==address]
                    if len(matches)>1:fail('当前区域管理 IP 已重复，不能自动匹配，请使用唯一记录编号')
                    if matches:
                        if mode=='create_only':fail('当前区域管理 IP 已存在，请选择更新模式，或移除重复的新增行')
                        if data.get('code'):fail('记录编号与已有管理 IP 不一致，请保留导出编号或将编号留空')
                        target=matches[0]['id'];expected=matches[0]['version'];data['code']=matches[0]['code']
                code=data.get('code') or L.catalog.generated_code(kind);data['code']=code
                if code in seen:fail('模板内编号重复')
                seen.add(code)
                if target:data['version']=expected
                if target:
                    previous=L.detail(conn,user,target,False)
                    if 'site_id' not in data:data['site_id']=previous['site_id']
                    if kind=='vpn' and 'remote_site_id' not in data:data['remote_site_id']=previous.get('remote_site_id')
                if kind=='device':
                    data['_import_order']=position
                    if data.get('uplink_device_id'):data['_uplink_reference']=data['uplink_device_id']
                data=resolve(conn,user,kind,data)
                if kind=='wifi' and data.get('auth_method') in ('不认证','密码','密码+MAC'):data['auth_method']={'不认证':'none','密码':'password','密码+MAC':'password_mac'}[data['auth_method']]
                if kind=='vpn':
                    modes={'指定 VLAN 网段':'selected','指定网段':'selected','selected':'selected','Any':'any','ANY':'any','any':'any','0.0.0.0/0':'any'}
                    for name in ('local_network_mode','remote_network_mode'):
                        if data.get(name) in modes:data[name]=modes[data[name]]
                original=dict(row(conn,'resources',target)) if target else {}
                if original:require(conn,user,'resource.import',original)
                require(conn,user,'resource.import',{**original,**data,'created_by_user_id':original.get('created_by_user_id',user['id'])})
                # Real validation, uniqueness and permissions in a rollback-only savepoint.
                tx=conn.begin_nested()
                try:
                    L.save(conn,user,kind,{k:v for k,v in data.items() if k not in ('_import_order','_uplink_reference')},target)
                    if kind=='device':tx.commit()
                finally:
                    if tx.is_active:tx.rollback()
            except Exception as exc:
                if hasattr(exc,'detail'):error=str(exc.detail)
                elif isinstance(exc,(ValueError,TypeError,csv.Error)):error='字段格式无效'
                elif isinstance(exc,sa.exc.IntegrityError):error='VLAN ID、SSID 或关联数据冲突'
                else:raise
            data=L.protect_secrets(conn,kind,data,encrypt=True)
            staged.append((line,jsonable(data),target,expected,error))
    finally:
        if device_tx is not None and device_tx.is_active:device_tx.rollback()
        wb.close()
    if not staged:fail('模板没有数据')
    template=T['import_templates'];code='ledger_'+kind
    version=template_version(kind)
    tpl=conn.execute(sa.select(template).where(template.c.template_code==code,template.c.template_version==version)).mappings().first()
    if not tpl:
        spec=json.dumps(cs,ensure_ascii=False)
        tpl=insert_row(conn,'import_templates',{'template_code':code,'template_version':version,'module_code':'resource','name':L.KINDS[kind][0],'file_format':'xlsx','field_spec_json':spec,'schema_sha256':digest(spec),'is_enabled':True},user['id'])
    sources=L.rows(conn,'data_sources',source_code='ledger_template')
    source=sources[0] if sources else insert_row(conn,'data_sources',{'source_code':'ledger_template','name':'五类台账模板','source_type':'standard_template','safe_config_json':'{}','is_enabled':True},user['id'])
    batch=insert_row(conn,'import_batches',{'source_id':source['id'],'requested_by_user_id':user['id'],'original_filename':filename.replace('\\','/').split('/')[-1][:256],
        'source_sha256':hashlib.sha256(raw).hexdigest(),'template_version':version,'batch_state':'review','template_id':tpl['id'],'import_mode':mode,
        'summary_json':json.dumps({'module':'resources','ledger_kind':kind,'rows':len(staged),'errors':sum(bool(v[4]) for v in staged)})},user['id'])
    for line,data,target,expected,error in staged:
        insert_row(conn,'import_rows',{'batch_id':batch['id'],'source_sheet':'数据','source_row_key':str(line),'entity_kind':'resources','source_external_key':str(line),
            'normalized_json':'{}','proposal_json':json.dumps(data,ensure_ascii=False),'validation_json':json.dumps({'error':error},ensure_ascii=False),'resolution_state':'conflict' if error else 'ready',
            'expected_version':expected,'target_entity_kind':'resources','target_entity_id':target},user['id'])
    audit(conn,user,'import.preview','import_batches',batch['id'],{'ledger_kind':kind,'rows':len(staged)})
    return batch


def commit(conn,user,batch,valid_only):
    require(conn,user,'resource.import')
    if batch['requested_by_user_id']!=user['id']:fail('只能确认自己的导入批次',403)
    if batch['batch_state']=='committed':return dict(batch)
    info=json.loads(batch['summary_json'])
    if batch['batch_state']!='review' or (info['errors'] and not valid_only):fail('批次状态无效或仍存在错误行',409)
    t=T['import_batches']
    if conn.execute(t.update().where(t.c.id==batch['id'],t.c.version==batch['version'],t.c.batch_state=='review').values(batch_state='committing',version=batch['version']+1)).rowcount!=1:fail('批次已被修改',409)
    imported=0
    batch_rows=L.rows(conn,'import_rows',batch_id=batch['id'])
    batch_rows.sort(key=lambda item:json.loads(item['proposal_json']).get('_import_order',int(item['source_row_key'])))
    for item in batch_rows:
        if item['resolution_state']!='ready':continue
        data=json.loads(item['proposal_json'])
        data=L.protect_secrets(conn,info['ledger_kind'],data,include_secret=True)
        if info['ledger_kind']=='device' and '_uplink_reference' in data:
            data['uplink_device_id']=data.pop('_uplink_reference')
            data=resolve(conn,user,'device',data)
        data.pop('_import_order',None)
        if item['target_entity_id']:data['version']=item['expected_version']
        original=dict(row(conn,'resources',item['target_entity_id']) or {})
        if original:require(conn,user,'resource.import',original)
        require(conn,user,'resource.import',{**original,**data,'created_by_user_id':original.get('created_by_user_id',user['id'])})
        obj=L.save(conn,user,info['ledger_kind'],data,item['target_entity_id']);imported+=1
        conn.execute(T['import_rows'].update().where(T['import_rows'].c.id==item['id']).values(resolution_state='committed',target_entity_id=obj['id']))
    if not imported:fail('没有可提交的正确行')
    conn.execute(t.update().where(t.c.id==batch['id']).values(batch_state='committed',committed_at=utcnow(),confirmed_by_user_id=user['id']))
    audit(conn,user,'import.commit','import_batches',batch['id'],{'rows':imported})
    return dict(row(conn,'import_batches',batch['id']))
