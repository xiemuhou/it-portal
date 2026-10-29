"""Versioned XLSX templates and staged atomic imports. No legacy source fetching."""
from io import BytesIO
from datetime import datetime
import hashlib
import json
import zipfile
import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .database import TABLES as T, SPECS, insert_row, row, utcnow, jsonable
from .security import require, fail, digest, audit
from . import catalog

VERSION='1.0'
MAX_BYTES=20*1024*1024


def business_key(name):
    names=[f['name'] for f in SPECS[name]['fields']]
    for key in ('resource_code','announcement_code','service_code','project_code','code','scope_code'):
        if key in names:return [key]
    return [k.strip() for k in SPECS[name]['unique'].split(';')[0].split(',') if k.strip()] or [SPECS[name]['pk']]


def reference_value(conn,target,key):
    obj=row(conn,target,key)
    keys=business_key(target) if target in catalog.MODULES else ['login_name'] if target=='users' else ['id']
    return obj[keys[0]] if obj and len(keys)==1 else key


def resolve_references(conn,name,data):
    result=dict(data)
    for f in catalog.fields(name):
        value=result.get(f['name'])
        if f['ref'] and value not in (None,'','__CLEAR__'):
            target=f['ref'].split('.')[0];table=T[target]
            keys=business_key(target) if target in catalog.MODULES else ['login_name'] if target=='users' else ['id']
            if len(keys)!=1:keys=['id']
            found=conn.execute(sa.select(table.c[SPECS[target]['pk']]).where(table.c[keys[0]]==str(value))).first()
            if not found and target=='sites':
                alias=T['site_aliases'];found=conn.execute(sa.select(alias.c.site_id).where(alias.c.alias_key==digest(str(value)))).first()
            if not found:fail(f['label']+'引用的业务编码不存在，请先导入关联对象')
            result[f['name']]=found[0]
    return result


def workbook(name,items=None,conn=None):
    fs=catalog.fields(name);wb=Workbook();sheet=wb.active;sheet.title='数据'
    sheet.append([f['header'] for f in fs])
    if items:
        for item in items:
            values=[]
            for f in fs:
                value=item.get(f['name'])
                if f['ref'] and value:value=reference_value(conn,f['ref'].split('.')[0],value)
                if isinstance(value,datetime):value=value.isoformat(timespec='milliseconds')+'Z'
                if hasattr(value,'as_tuple'):value=str(value)
                values.append(value)
            sheet.append(values)
    for r in sheet:
        for cell in r:
            if isinstance(cell.value,str):cell.data_type='s'
    for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='2764E7')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=28
    meta=wb.create_sheet('模板信息');meta.append(['模板代码',name]);meta.append(['模板版本',VERSION]);meta.append(['用途','标准数据模板，可回导；汇总报表不可回导'])
    guide=wb.create_sheet('填写说明');guide.append(['列名','必填','类型','引用/选项','规则'])
    for f in fs:guide.append([f['header'],'否' if f['nullable'] or f['auto_generated'] else '是',f['type'],f['ref'] or ' / '.join(f['enum']),('新增留空由系统生成；更新已有记录时保留导出编号。' if f['auto_generated'] else '')+f['description']])
    guide.append(['更新规则','','','','更新模式下留空保留原值；__CLEAR__ 明确清空；先导入基础和主记录'])
    guide.append(['时间规则','','','','日期 YYYY-MM-DD；带时间的值使用带时区 ISO 格式，如 2026-10-16T09:00:00+08:00'])
    guide.append(['发布规则','','','','公告/成果导入仅成草稿，不能以模板直接发布'])
    for col in guide.columns:guide.column_dimensions[col[0].column_letter].width=38
    out=BytesIO();wb.save(out);return out.getvalue()


def ensure_template(conn,name,user):
    t=T['import_templates'];source=T['data_sources']
    template=conn.execute(sa.select(t).where(t.c.template_code==name,t.c.template_version==VERSION)).mappings().first()
    if not template:
        spec=json.dumps(catalog.fields(name),ensure_ascii=False,sort_keys=True)
        template=insert_row(conn,'import_templates',{'template_code':name,'template_version':VERSION,'module_code':catalog.module(name),
            'name':SPECS[name]['title'],'file_format':'xlsx','field_spec_json':spec,'schema_sha256':digest(spec),'is_enabled':True},user['id'])
    data_source=conn.execute(sa.select(source).where(source.c.source_code=='standard_template')).mappings().first()
    if not data_source:data_source=insert_row(conn,'data_sources',{'source_code':'standard_template','name':'标准模板人工导入','source_type':'standard_template','safe_config_json':'{}','is_enabled':True},user['id'])
    return template,data_source


def preview(conn,user,name,raw,filename,mode='create_only'):
    mod=catalog.module(name);require(conn,user,mod+'.import')
    if mode not in ('create_only','upsert'):fail('导入模式无效')
    if not filename.lower().endswith('.xlsx') or len(raw)>MAX_BYTES:fail('仅支持 20 MB 以内的标准 XLSX 文件')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024 or any('vbaproject' in i.filename.lower() for i in z.infolist()):fail('文件内容过大或包含宏')
        wb=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except (zipfile.BadZipFile,ValueError,KeyError):fail('无法读取 XLSX 文件')
    try:
        if set(wb.sheetnames)!={'数据','模板信息','填写说明'}:fail('请下载本模块标准模板，不要修改工作表结构')
        meta=dict(wb['模板信息'].iter_rows(values_only=True))
        if meta.get('模板代码')!=name or str(meta.get('模板版本'))!=VERSION:fail('模板类型或版本不匹配，请重新下载')
        fs=catalog.fields(name);it=iter(wb['数据'].iter_rows());headers=next(it,[])
        if [c.value for c in headers]!=[f['header'] for f in fs]:fail('表头与标准模板不一致')
        staged=[];seen=set();keys=business_key(name)
        for index,cells in enumerate(it,2):
            if index>5001:fail('单次最多导入 5000 行，请拆分批次')
            if not any(c.value is not None for c in cells):continue
            data={f['name']:c.value for f,c in zip(fs,cells) if c.value is not None}
            # Never retain unknown fields, formulas or credentials in staging.
            error=None;proposal=None;existing=None
            try:
                if any(c.data_type=='f' for c in cells):fail('业务字段不能使用公式')
                data=resolve_references(conn,name,data)
                for key in keys:
                    if key in catalog.AUTO_CODES and not data.get(key):
                        data[key]=catalog.generated_code(name)
                if any(k not in data for k in keys):fail('缺少用于匹配的稳定业务编号')
                identity=json.dumps([str(data[k]) for k in keys])
                if identity in seen:fail('文件中存在重复业务编号')
                seen.add(identity)
                table=T[name]
                existing=conn.execute(sa.select(table).where(*[table.c[k]==data[k] for k in keys])).mappings().first()
                if existing and mode=='create_only':fail('业务编号已存在，若需更新请选择更新模式')
                data={k:None if v=='__CLEAR__' else v for k,v in data.items()}
                proposal=catalog.validate(conn,user,name,data,existing,True)
                permission=mod+'.'+('edit' if existing or mod!='resource' else 'create')
                require(conn,user,permission,catalog.contextual(conn,name,{**dict(existing or {}),**proposal,'created_by_user_id':existing.get('created_by_user_id') if existing else user['id']}))
            except Exception as exc:
                if hasattr(exc,'detail'):error=str(exc.detail)
                else:raise
            staged.append({'line':index,'proposal':jsonable(proposal) if proposal else {},'target':dict(existing) if existing else None,'error':error})
    finally:wb.close()
    if not staged:fail('模板中没有数据')
    template,source=ensure_template(conn,name,user)
    batch=insert_row(conn,'import_batches',{'source_id':source['id'],'requested_by_user_id':user['id'],
        'original_filename':PathName(filename),'source_sha256':hashlib.sha256(raw).hexdigest(),'template_version':VERSION,'batch_state':'review',
        'template_id':template['id'],'import_mode':mode,'summary_json':json.dumps({'module':name,'rows':len(staged),'errors':sum(bool(v['error']) for v in staged)})},user['id'])
    for r in staged:
        target=r['target'];version=target.get('version') if target else None
        if target and 'version' not in target and 'resource_id' in target:version=row(conn,'resources',target['resource_id'])['version']
        insert_row(conn,'import_rows',{'batch_id':batch['id'],'source_sheet':'数据','source_row_key':str(r['line']),
            'entity_kind':name,'source_external_key':str(r['line']), 'normalized_json':'{}','proposal_json':json.dumps(r['proposal'],ensure_ascii=False),
            'validation_json':json.dumps({'error':r['error']},ensure_ascii=False),'resolution_state':'conflict' if r['error'] else 'ready',
            'expected_version':version,'target_entity_kind':name,'target_entity_id':target[SPECS[name]['pk']] if target else None},user['id'])
    audit(conn,user,'import.preview','import_batches',batch['id'],{'module':name,'rows':len(staged)})
    return batch


def PathName(value):
    return str(value).replace('\\','/').rsplit('/',1)[-1][:256]


def commit(conn,user,key,valid_only=False):
    batch=row(conn,'import_batches',key)
    if not batch:fail('导入批次不存在',404)
    if json.loads(batch['summary_json']).get('ledger_kind'):
        from . import ledger_transfers
        return ledger_transfers.commit(conn,user,batch,valid_only)
    info=json.loads(batch['summary_json']);name=info['module'];mod=catalog.module(name)
    require(conn,user,mod+'.import')
    if batch['requested_by_user_id']!=user['id']:fail('只能确认自己上传的导入批次',403)
    if batch['batch_state']=='committed':return batch
    if batch['batch_state']!='review':fail('此批次不可提交',409)
    if info['errors'] and not valid_only:fail('存在错误行，默认整批不提交；可显式选择仅提交正确行',409)
    t=T['import_batches']
    if conn.execute(t.update().where(t.c.id==key,t.c.version==batch['version'],t.c.batch_state=='review').values(batch_state='committing',version=batch['version']+1)).rowcount!=1:fail('批次已被确认',409)
    rows=T['import_rows'];result=[]
    # Reuse this connection for writes only after the SELECT cursor is drained.
    staged_rows=conn.execute(sa.select(rows).where(rows.c.batch_id==key).order_by(rows.c.source_row_key)).mappings().all()
    for staged in staged_rows:
        if staged['resolution_state']!='ready':continue
        proposal=json.loads(staged['proposal_json'])
        # Derived fields cannot be accepted from import; validation recomputes them.
        proposal={k:v for k,v in proposal.items() if k not in catalog.READONLY}
        obj=catalog.save(conn,user,name,proposal,staged['target_entity_id'],True,staged['expected_version'])
        result.append(obj[SPECS[name]['pk']])
        conn.execute(rows.update().where(rows.c.id==staged['id']).values(resolution_state='committed',target_entity_id=obj[SPECS[name]['pk']]))
    if not result:fail('没有可提交的正确行')
    conn.execute(t.update().where(t.c.id==key).values(batch_state='committed',committed_at=utcnow(),confirmed_by_user_id=user['id']))
    audit(conn,user,'import.commit','import_batches',key,{'committed':len(result),'skipped':info['errors']})
    return dict(row(conn,'import_batches',key))
