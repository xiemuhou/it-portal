"""Static IPv4 inventory for explicitly enabled VLAN networks and computer assets."""
from datetime import date,datetime
import ipaddress
import json
import re
from io import BytesIO

import sqlalchemy as sa
from fastapi import Request,UploadFile,File
from fastapi.responses import Response
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill

from .database import TABLES as T, insert_row, row, utcnow, jsonable, uid
from .security import require, allow, fail, audit
from . import catalog, ledger


OCCUPANCY = {'available','occupied','reserved','disabled','conflict'}
VERIFICATION = {'pending','confirmed'}
ASSET_STATES = {'in_use','standby','repair','retired'}
ADAPTER_TYPES = {'wired','wifi','virtual','other'}
COMPUTER_DEVICE_TYPES = {'笔记本','台式机'}
IP_HEADERS=['区域','VLAN ID','网段','IP 地址','归属类型','部门','使用人','占用设备','MAC地址','位置','备注','状态','确认状态']
COMPUTER_HEADERS=['资产编号','设备类型','验收日期','归属类型','区域','部门','使用人','位置','操作系统','状态','网卡类型','MAC','关联 IP','配件类型','配件备注','电脑备注']

OCCUPANCY_CN={'空闲':'available','已占用':'occupied','预留':'reserved','停用':'disabled','冲突':'conflict'}
VERIFICATION_CN={'待确认':'pending','已确认':'confirmed'}
ASSET_CN={'在用':'in_use','备用':'standby','维修':'repair','报废':'retired'}
ADAPTER_CN={'有线':'wired','有线网卡':'wired','无线':'wifi','无线网卡':'wifi','虚拟':'virtual','虚拟网卡':'virtual','其他':'other','其他网卡':'other'}

def _download(data,filename):
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition':f"attachment; filename*=UTF-8''{filename}"})

def _cell(value):
    return str(value).strip() if value not in (None,'') else ''

def _named_id(values,name,label,required=False):
    text=_cell(name)
    if not text and not required:return None
    matches=[key for key,value in values.items() if value==text]
    if len(matches)!=1:fail(f'{label}“{text}”不存在或不唯一')
    return matches[0]

def _xlsx(headers,rows=()):
    wb=Workbook();ws=wb.active;ws.title='数据';ws.append(headers)
    for row_values in rows:ws.append(row_values)
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='287DDD')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=22
    guide=wb.create_sheet('填写说明');guide.append(['说明','请勿修改“数据”工作表表头；同一电脑有多块网卡、多个 IP 或多个配件时，可重复资产编号分多行填写。'])
    out=BytesIO();wb.save(out);return out.getvalue()

def _read_xlsx(raw,headers):
    if len(raw)>20*1024*1024:fail('文件不能超过 20 MB')
    try:wb=load_workbook(BytesIO(raw),read_only=True,data_only=True)
    except Exception:fail('无法读取 XLSX 文件')
    ws=wb['数据'] if '数据' in wb.sheetnames else wb.active;values=ws.iter_rows(values_only=True)
    if list(next(values,()))!=headers:fail('表头与标准模板不一致，请重新下载模板')
    result=[dict(zip(headers,row)) for row in values if any(v not in (None,'') for v in row)];wb.close()
    if not result:fail('模板中没有数据')
    return result


def _text(value,label,limit,required=False):
    if value is None:value=''
    if not isinstance(value,str):fail(label+'格式无效')
    value=value.strip()
    if required and not value:fail('请填写'+label)
    if '\x00' in value or len(value)>limit:fail(label+'内容过长或无效')
    return value or None


def _month(value):
    if value in (None,''):return None
    if isinstance(value,(date,datetime)):result=date(value.year,value.month,1)
    else:
        text=str(value).strip().replace('/','-').replace('.','-')
        match=re.fullmatch(r'(\d{4})-(\d{1,2})(?:-\d{1,2})?',text)
        if not match:fail('验收日期格式无效，请使用 YYYY-MM')
        try:result=date(int(match.group(1)),int(match.group(2)),1)
        except ValueError:fail('验收日期格式无效，请使用 YYYY-MM')
    current=date.today().replace(day=1)
    if result>current:fail('验收日期不能晚于当前月份')
    return result


def _mac(value,required=False):
    raw=re.sub(r'[^0-9A-Fa-f]','',str(value or ''))
    if not raw and not required:return None
    if len(raw)!=12:fail('MAC 地址须为 12 位十六进制字符')
    return ':'.join(raw[i:i+2] for i in range(0,12,2)).upper()


def _lookups(conn):
    departments={v['id']:v['name'] for v in conn.execute(sa.select(T['departments'].c.id,T['departments'].c.name).where(T['departments'].c.is_archived==False)).mappings()}
    users={v['id']:v['display_name'] for v in conn.execute(sa.select(T['users'].c.id,T['users'].c.display_name).where(T['users'].c.is_archived==False)).mappings()}
    sites={v['id']:v['name'] for v in conn.execute(sa.select(T['sites'].c.id,T['sites'].c.name).where(T['sites'].c.is_archived==False)).mappings()}
    return departments,users,sites


def _computer(conn,user,key):
    value=row(conn,'computer_assets',key)
    if not value or value['is_archived']:fail('电脑设备不存在',404)
    require(conn,user,'resource.view',value)
    result=dict(value)
    if result.get('acceptance_month'):
        current=date.today();accepted=result['acceptance_month']
        months=max(0,(current.year-accepted.year)*12+current.month-accepted.month)
        result['usage_months']=months;result['usage_text']=f'{months//12}年{months%12}个月';result['is_over_five_years']=months>60
    else:
        result['usage_months']=None;result['usage_text']=None;result['is_over_five_years']=False
    result['adapters']=[dict(v) for v in conn.execute(sa.select(T['computer_network_adapters']).where(T['computer_network_adapters'].c.computer_id==key).order_by(T['computer_network_adapters'].c.is_primary.desc(),T['computer_network_adapters'].c.adapter_name)).mappings()]
    result['ip_addresses']=[dict(v) for v in conn.execute(sa.select(
        T['static_ip_records'].c.id,T['static_ip_records'].c.ip_address,T['static_ip_records'].c.occupancy_status,
        T['static_ip_records'].c.verification_status,T['static_ip_records'].c.mac_address
    ).where(T['static_ip_records'].c.computer_id==key,T['static_ip_records'].c.occupancy_status!='available')).mappings()]
    result['ip_addresses'].sort(key=lambda value:ipaddress.ip_address(value['ip_address']))
    result['accessories']=[dict(v) for v in conn.execute(sa.select(T['computer_accessories']).where(T['computer_accessories'].c.computer_id==key).order_by(T['computer_accessories'].c.accessory_type,T['computer_accessories'].c.id)).mappings()]
    if 'consumable_movements' in T:
        movements,items=T['consumable_movements'],T['consumable_items']
        result['consumable_issues']=[dict(v) for v in conn.execute(sa.select(movements.c.id,movements.c.occurred_on,movements.c.quantity,
            movements.c.unit_snapshot,movements.c.purpose,movements.c.notes,items.c.name,items.c.specification).join(items,items.c.id==movements.c.item_id).where(
            movements.c.computer_id==key,movements.c.movement_type=='outbound',movements.c.is_reversed==False).order_by(movements.c.occurred_on.desc(),movements.c.created_at.desc())).mappings()]
    departments,users,sites=_lookups(conn)
    result.update(site_name=sites.get(result['site_id']),department_name=departments.get(result['department_id']),display_user=users.get(result['user_id']) or result.get('user_name'))
    return result

def _save_accessories(conn,computer_id,items):
    if not isinstance(items,list) or len(items)>50:fail('配件列表格式无效或数量过多')
    table=T['computer_accessories'];conn.execute(table.delete().where(table.c.computer_id==computer_id));now=utcnow()
    for item in items:
        if not isinstance(item,dict):fail('配件信息格式无效')
        conn.execute(table.insert().values(id=uid(),created_at=now,updated_at=now,computer_id=computer_id,
            accessory_type=_text(item.get('accessory_type'),'配件类型',64,True),brand_model=_text(item.get('brand_model'),'品牌型号',160),
            serial_number=_text(item.get('serial_number'),'序列号',160),notes=_text(item.get('notes'),'配件备注',512)))

def _sync_computer_ips(conn,user,computer,assignments):
    if assignments is None:return
    if not isinstance(assignments,list) or len(assignments)>20:fail('静态 IP 列表格式无效或数量过多')
    desired=[]
    for raw in assignments:
        item=raw if isinstance(raw,dict) else {'ip_address':raw,'adapter_mac':None}
        try:value=str(ipaddress.ip_address(str(item.get('ip_address') or '').strip()))
        except ValueError:fail('静态 IP 地址无效：'+str(item.get('ip_address') or raw))
        if value not in [v['ip_address'] for v in desired]:desired.append({'ip_address':value,'adapter_mac':_mac(item.get('adapter_mac'))})
    options=[v for v in _managed_options(conn,user) if v['site_id']==computer['site_id']]
    adapters=conn.execute(sa.select(T['computer_network_adapters']).where(T['computer_network_adapters'].c.computer_id==computer['id'],T['computer_network_adapters'].c.enabled==True).order_by(T['computer_network_adapters'].c.is_primary.desc())).mappings().all()
    adapter_by_mac={v['mac_address']:v for v in adapters};t=T['static_ip_records'];now=utcnow()
    current=conn.execute(sa.select(t).where(t.c.computer_id==computer['id'],t.c.occupancy_status!='available')).mappings().all()
    for assignment in desired:
        address=assignment['ip_address'];adapter=adapter_by_mac.get(assignment['adapter_mac'])
        if not adapter:fail('请为 IP '+address+' 选择属于此电脑的网卡')
        matches=[]
        for option in options:
            net=ipaddress.ip_network(option['cidr'],strict=True)
            if ipaddress.ip_address(address) in net and address not in (str(net.network_address),str(net.broadcast_address),option.get('gateway')):matches.append(option)
        if len(matches)!=1:fail('IP '+address+' 未唯一匹配当前区域已启用静态 IP 管理的网段')
        option=matches[0];old=conn.execute(sa.select(t).where(t.c.vlan_id==option['vlan_id'],t.c.network_id==option['network_id'],t.c.ip_address==address)).mappings().first()
        values={'occupancy_status':'occupied','verification_status':'confirmed','assignment_type':computer.get('assignment_type') or 'personal','department_id':computer['department_id'],'department_name_snapshot':None,'user_id':computer['user_id'],'user_name':computer['user_name'],'device_type':computer['device_type'],'computer_id':computer['id'],'adapter_id':adapter and adapter['id'],'mac_address':adapter and adapter['mac_address'],'location':computer['location'],'verified_at':now,'verified_by_user_id':user['id'],'updated_at':now,'updated_by_user_id':user['id']}
        if old:
            if old['occupancy_status']!='available':
                if old.get('computer_id') not in (None,computer['id']):fail('IP '+address+' 已被其他电脑占用',409)
                if old.get('computer_id') is None and old.get('mac_address')!=adapter['mac_address']:fail('IP '+address+' 的 MAC 与所选电脑网卡不一致，不能自动关联',409)
                values={'computer_id':computer['id'],'adapter_id':adapter['id'],'mac_address':adapter['mac_address'],'updated_at':now,'updated_by_user_id':user['id']}
            conn.execute(t.update().where(t.c.id==old['id']).values(**values,version=old['version']+1));record_id=old['id']
        else:record_id=insert_row(conn,'static_ip_records',{'vlan_id':option['vlan_id'],'network_id':option['network_id'],'ip_address':address,'notes':None,**values},user['id'])['id']
        saved=dict(row(conn,'static_ip_records',record_id));insert_row(conn,'static_ip_history',{'ip_record_id':record_id,'actor_user_id':user['id'],'action_code':'computer_linked','snapshot_json':json.dumps(jsonable(saved),ensure_ascii=False)})
    for old in current:
        if old['ip_address'] in [v['ip_address'] for v in desired]:continue
        conn.execute(t.update().where(t.c.id==old['id']).values(occupancy_status='available',verification_status='pending',department_id=None,department_name_snapshot=None,user_id=None,user_name=None,device_type=None,computer_id=None,adapter_id=None,mac_address=None,location=None,verified_at=None,verified_by_user_id=None,version=old['version']+1,updated_at=now,updated_by_user_id=user['id']))
        saved=dict(row(conn,'static_ip_records',old['id']));insert_row(conn,'static_ip_history',{'ip_record_id':old['id'],'actor_user_id':user['id'],'action_code':'computer_unlinked','snapshot_json':json.dumps(jsonable(saved),ensure_ascii=False)})

def _record_computer_history(conn,user,computer_id,action,summary):
    snapshot=_computer(conn,user,computer_id)
    insert_row(conn,'computer_change_history',{'computer_id':computer_id,'actor_user_id':user['id'],'action_code':action,'change_summary':summary,'snapshot_json':json.dumps(jsonable(snapshot),ensure_ascii=False)})


def _mac_asset_candidate(conn,record,site_id):
    """Return the only active computer adapter matching an unlinked IP MAC."""
    if not record.get('mac_address'):return None
    adapters=T['computer_network_adapters'];computers=T['computer_assets']
    matches=conn.execute(sa.select(adapters,computers.c.asset_code,computers.c.device_type,computers.c.site_id,
        computers.c.user_id,computers.c.department_id,computers.c.assignment_type,computers.c.location,
        computers.c.version.label('computer_version')).select_from(adapters.join(computers,computers.c.id==adapters.c.computer_id)).where(
        adapters.c.mac_address==record['mac_address'],adapters.c.enabled==True,computers.c.is_archived==False,computers.c.site_id==site_id)).mappings().all()
    return dict(matches[0]) if len(matches)==1 else None


def _managed_network(conn,user,key):
    if ':' not in key:fail('请选择静态 IP 受管网段')
    vlan_id,network_id=key.split(':',1)
    vlan=ledger.detail(conn,user,vlan_id,False)
    network=next((n for n in vlan.get('networks',[]) if n['network_id']==network_id),None)
    if not network or not network.get('ipam_enabled'):fail('该网段未启用静态 IP 管理',404)
    try:net=ipaddress.ip_network(network['cidr'],strict=True)
    except ValueError:fail('网段配置无效')
    if net.version!=4 or net.num_addresses>65536:fail('静态 IP 管理当前支持最多 65536 个地址的 IPv4 网段')
    return vlan,network,net


def _managed_options(conn,user):
    result=[]
    for vlan in ledger.list_items(conn,user,'vlan'):
        for network in vlan.get('networks',[]):
            if network.get('ipam_enabled'):
                result.append({'id':ledger.network_key(vlan['id'],network['network_id']),'vlan_id':vlan['id'],'network_id':network['network_id'],
                    'site_id':vlan['site_id'],'site_name':vlan['site_name'],'vlan_number':vlan['vlan_number'],**network})
    return result


def _save_adapters(conn,user,computer_id,items):
    if not isinstance(items,list) or len(items)>20:fail('网卡列表格式无效或数量过多')
    table=T['computer_network_adapters'];existing={v['id']:dict(v) for v in conn.execute(sa.select(table).where(table.c.computer_id==computer_id)).mappings()}
    seen_ids=set();seen_macs=set();now=utcnow()
    for item in items:
        if not isinstance(item,dict):fail('网卡信息格式无效')
        key=item.get('id');mac=_mac(item.get('mac_address'),True)
        if key and (key not in existing or key in seen_ids):fail('网卡记录已变化，请刷新')
        if mac in seen_macs:fail('同一电脑不能重复登记 MAC 地址')
        owner=conn.execute(sa.select(table.c.id).where(table.c.mac_address==mac,table.c.id!=(key or ''))).first()
        if owner:fail('MAC 地址已关联其他电脑：'+mac,409)
        adapter_type=item.get('adapter_type') or 'wired'
        if adapter_type not in ADAPTER_TYPES:fail('网卡类型无效')
        previous=existing.get(key) if key else None
        values={'adapter_name':_text(item.get('adapter_name'),'网卡名称',128) or (previous and previous['adapter_name']) or {'wired':'有线网卡','wifi':'无线网卡','virtual':'虚拟网卡','other':'其他网卡'}[adapter_type],
            'adapter_type':adapter_type,'mac_address':mac,
            'is_primary':bool(item['is_primary']) if 'is_primary' in item else bool(previous and previous['is_primary']),
            'enabled':item.get('enabled') is not False if 'enabled' in item else bool(previous['enabled']) if previous else True,
            'notes':_text(item.get('notes'),'网卡备注',512) if 'notes' in item else previous and previous['notes']}
        if key:
            conn.execute(table.update().where(table.c.id==key).values(**values,updated_at=now));seen_ids.add(key)
        else:
            insert_row(conn,'computer_network_adapters',{'computer_id':computer_id,**values});
    for key,old in existing.items():
        if key in seen_ids:continue
        used=conn.execute(sa.select(T['static_ip_records'].c.id).where(T['static_ip_records'].c.adapter_id==key,T['static_ip_records'].c.occupancy_status!='available').limit(1)).first()
        if used:conn.execute(table.update().where(table.c.id==key).values(enabled=False,updated_at=now))
        else:conn.execute(table.delete().where(table.c.id==key))


def register(app,engine,out,current_user):
    @app.get('/api/ipam/options')
    def options(request:Request):
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.view')
            departments,users,sites=_lookups(conn)
            computers=[]
            query=sa.select(T['computer_assets']).where(T['computer_assets'].c.is_archived==False,catalog.scope_expression(conn,user,'resource.view',T['computer_assets']))
            # Materialize the outer result before _computer issues more statements on
            # the same connection. SQL Server ODBC does not enable MARS by default.
            for value in conn.execute(query.order_by(T['computer_assets'].c.device_name)).mappings().all():
                computers.append(_computer(conn,user,value['id']))
            user_options=[]
            for value in conn.execute(sa.select(T['users'].c.id,T['users'].c.display_name,T['users'].c.department_id).where(T['users'].c.is_archived==False).order_by(T['users'].c.display_name)).mappings():
                user_options.append({'id':value['id'],'name':value['display_name'],'department_id':value['department_id'],'department_name':departments.get(value['department_id'])})
            return out({'networks':_managed_options(conn,user),'departments':[{'id':k,'name':v} for k,v in departments.items()],
                'users':user_options,'sites':[{'id':k,'name':v} for k,v in sites.items()],'computers':computers})

    @app.get('/api/computers')
    def computers(request:Request,page:int=1,page_size:int=10,q:str='',site_id:str=''):
        if page_size not in (10,20,50):fail('每页条数仅支持 10、20、50')
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.view');t=T['computer_assets']
            query=sa.select(t.c.id).where(t.c.is_archived==False,catalog.scope_expression(conn,user,'resource.view',t))
            if site_id:query=query.where(t.c.site_id==site_id)
            if q:query=query.where(sa.or_(t.c.asset_code.contains(q,autoescape=True),t.c.device_name.contains(q,autoescape=True),t.c.serial_number.contains(q,autoescape=True),t.c.user_name.contains(q,autoescape=True)))
            total=conn.execute(sa.select(sa.func.count()).select_from(query.subquery())).scalar_one();page=max(1,page)
            keys=conn.execute(query.order_by(t.c.device_name,t.c.id).offset((page-1)*page_size).limit(page_size)).scalars().all()
            return out({'items':[_computer(conn,user,k) for k in keys],'total':total,'page':page,'page_size':page_size})

    @app.get('/api/computers/{key}')
    def computer(request:Request,key:str):
        with engine.begin() as conn:return out(_computer(conn,current_user(conn,request),key))

    @app.get('/api/computers/{key}/history')
    def computer_history(request:Request,key:str):
        with engine.begin() as conn:
            user=current_user(conn,request);_computer(conn,user,key);h=T['computer_change_history'];u=T['users']
            values=conn.execute(sa.select(h,u.c.display_name.label('actor_name')).join(u,u.c.id==h.c.actor_user_id).where(h.c.computer_id==key).order_by(h.c.created_at.desc())).mappings().all()
            return out({'items':[dict(v) for v in values]})

    def save_computer(conn,user,body,key=None):
        old=row(conn,'computer_assets',key) if key else None
        before=_computer(conn,user,key) if old else None
        permission='resource.edit' if old else 'resource.create'
        if old:
            require(conn,user,permission,old)
            if body.get('version')!=old['version']:fail('电脑设备资料已变化，请刷新',409)
        site=row(conn,'sites',body.get('site_id'));assignment_type=body.get('assignment_type') or 'personal'
        if assignment_type not in ('personal','shared'):fail('电脑归属类型无效')
        linked_user=row(conn,'users',body.get('user_id')) if body.get('user_id') else None
        if not site or site['is_archived']:fail('请选择有效区域')
        if assignment_type=='personal' and (not linked_user or linked_user['is_archived']):fail('个人使用设备必须选择系统人员；部门将自动带出')
        if assignment_type=='shared':linked_user=None
        department=row(conn,'departments',linked_user['department_id']) if linked_user else None
        if linked_user and (not department or department['is_archived']):fail('所选人员未关联有效部门')
        target={'site_id':site['id'],'created_by_user_id':old['created_by_user_id'] if old else user['id']};require(conn,user,permission,target)
        status=body.get('asset_status') or 'in_use'
        if status not in ASSET_STATES:fail('设备状态无效')
        asset_code=_text(body.get('asset_code'),'资产编号',64) or (old and old['asset_code']) or ('PC-'+uid().replace('-','')[-10:].upper())
        device_type=_text(body.get('device_type'),'设备类型',64,True)
        if device_type not in COMPUTER_DEVICE_TYPES:fail('电脑设备类型仅支持笔记本或台式机')
        values={'asset_code':asset_code,
            'device_name':_text(body.get('device_name'),'设备名称',160) or (old and old['device_name']) or asset_code,'device_type':device_type,
            'brand':_text(body.get('brand'),'品牌',128),'model':_text(body.get('model'),'型号',160),'serial_number':_text(body.get('serial_number'),'序列号',160),
            'acceptance_month':_month(body.get('acceptance_month')) if 'acceptance_month' in body else old and old.get('acceptance_month'),
            'site_id':site['id'],'department_id':department['id'] if department else None,'user_id':linked_user['id'] if linked_user else None,
            'user_name':None,
            'location':_text(body.get('location'),'位置',256),'operating_system':_text(body.get('operating_system'),'操作系统',160),
            'assignment_type':assignment_type,'asset_status':status,'notes':_text(body.get('notes'),'备注',10000)}
        t=T['computer_assets'];now=utcnow()
        if old:
            result=conn.execute(t.update().where(t.c.id==key,t.c.version==old['version']).values(**values,version=old['version']+1,updated_at=now,updated_by_user_id=user['id']))
            if result.rowcount!=1:fail('电脑设备资料已变化，请刷新',409)
        else:key=insert_row(conn,'computer_assets',values,user['id'])['id']
        _save_adapters(conn,user,key,body.get('adapters',[]));_save_accessories(conn,key,body.get('accessories',[]))
        assignments=body.get('ip_assignments')
        if assignments is None and 'ip_addresses' in body:
            first_mac=next((v.get('mac_address') for v in body.get('adapters',[]) if v.get('mac_address')),None)
            assignments=[{'ip_address':v,'adapter_mac':first_mac} for v in body.get('ip_addresses',[])]
        _sync_computer_ips(conn,user,dict(row(conn,'computer_assets',key)),assignments)
        after=_computer(conn,user,key);changes=[]
        if not before:changes.append('建立电脑台账')
        else:
            if before.get('display_user')!=after.get('display_user'):changes.append(f"使用人：{before.get('display_user') or '未分配'} → {after.get('display_user') or '未分配'}")
            if before.get('department_name')!=after.get('department_name'):changes.append(f"部门：{before.get('department_name') or '未填写'} → {after.get('department_name') or '未填写'}")
            if before.get('location')!=after.get('location'):changes.append(f"位置：{before.get('location') or '未填写'} → {after.get('location') or '未填写'}")
            if before.get('acceptance_month')!=after.get('acceptance_month'):changes.append(f"验收日期：{str(before.get('acceptance_month') or '未填写')[:7]} → {str(after.get('acceptance_month') or '未填写')[:7]}")
            old_parts='、'.join(v['accessory_type']+(('('+v['notes']+')') if v.get('notes') else '') for v in before['accessories']) or '无'
            new_parts='、'.join(v['accessory_type']+(('('+v['notes']+')') if v.get('notes') else '') for v in after['accessories']) or '无'
            if old_parts!=new_parts:changes.append(f'配件：{old_parts} → {new_parts}')
            old_ips='、'.join(v['ip_address']+'('+str(v.get('mac_address') or '未关联网卡')+')' for v in before['ip_addresses']) or '无'
            new_ips='、'.join(v['ip_address']+'('+str(v.get('mac_address') or '未关联网卡')+')' for v in after['ip_addresses']) or '无'
            if old_ips!=new_ips:changes.append(f'IP/网卡：{old_ips} → {new_ips}')
        _record_computer_history(conn,user,key,'created' if not old else 'updated','、'.join(changes) or '电脑资料更新');audit(conn,user,'computer.save','computer_assets',key)
        return _computer(conn,user,key)

    @app.get('/api/computer-transfers/template')
    def computer_template(request:Request):
        with engine.begin() as conn:require(conn,current_user(conn,request),'resource.import')
        return _download(_xlsx(COMPUTER_HEADERS),'computer-template.xlsx')

    @app.get('/api/computer-transfers/export')
    def computer_export(request:Request):
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.export');t=T['computer_assets']
            keys=conn.execute(sa.select(t.c.id).where(t.c.is_archived==False,catalog.scope_expression(conn,user,'resource.export',t)).order_by(t.c.asset_code)).scalars().all()
            rows=[];adapter_names={v:k for k,v in ADAPTER_CN.items() if '网卡' not in k};asset_names={v:k for k,v in ASSET_CN.items()}
            for key in keys:
                item=_computer(conn,user,key);slots=[]
                for adapter in item['adapters']:
                    ips=[v['ip_address'] for v in item['ip_addresses'] if v.get('mac_address')==adapter['mac_address']]
                    for address in ips or ['']:slots.append((adapter,address))
                linked={address for _,address in slots if address}
                slots.extend((None,v['ip_address']) for v in item['ip_addresses'] if v['ip_address'] not in linked)
                count=max(1,len(slots),len(item['accessories']))
                for index in range(count):
                    adapter,address=slots[index] if index<len(slots) else (None,'');part=item['accessories'][index] if index<len(item['accessories']) else None
                    rows.append([item['asset_code'],item['device_type'],item['acceptance_month'].strftime('%Y-%m') if item.get('acceptance_month') else '','公共设备' if item['assignment_type']=='shared' else '个人使用',item['site_name'],item['department_name'] or '',item['display_user'] or '',item['location'] or '',item['operating_system'] or '',asset_names.get(item['asset_status'],item['asset_status']),
                        adapter_names.get(adapter['adapter_type'],adapter['adapter_type']) if adapter else '',adapter['mac_address'] if adapter else '',address,
                        part['accessory_type'] if part else '',part['notes'] if part else '',item['notes'] or ''])
            audit(conn,user,'computer.export','computer_assets',None,{'count':len(keys)})
        return _download(_xlsx(COMPUTER_HEADERS,rows),'computer-assets.xlsx')

    @app.post('/api/computer-transfers/import')
    async def computer_import(request:Request,file:UploadFile=File(...)):
        data=_read_xlsx(await file.read(),COMPUTER_HEADERS);groups={}
        for line,item in enumerate(data,2):
            code=_cell(item['资产编号'])
            if not code:fail(f'第 {line} 行缺少资产编号')
            groups.setdefault(code,[]).append((line,item))
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.import');departments,users,sites=_lookups(conn);saved=[]
            existing={v['asset_code']:dict(v) for v in conn.execute(sa.select(T['computer_assets']).where(T['computer_assets'].c.is_archived==False)).mappings()}
            for code,group in groups.items():
                first=group[0][1];old=existing.get(code);site_id=_named_id(sites,first['区域'],'区域',True)
                assignment_type={'个人使用':'personal','公共设备':'shared'}.get(_cell(first['归属类型']),_cell(first['归属类型']) or 'personal')
                if assignment_type not in ('personal','shared'):fail(f'电脑 {code} 的归属类型无效')
                user_name=_cell(first['使用人']);user_id=_named_id(users,user_name,'使用人',assignment_type=='personal') if user_name or assignment_type=='personal' else None
                linked_user=row(conn,'users',user_id) if user_id else None
                if linked_user and _cell(first['部门']) and departments.get(linked_user['department_id'])!=_cell(first['部门']):fail(f'电脑 {code} 的部门与系统人员资料不一致')
                adapters=[];accessories=[];assignments=[];seen_macs=set();seen_parts=set()
                old_adapters={v['mac_address']:dict(v) for v in conn.execute(sa.select(T['computer_network_adapters']).where(T['computer_network_adapters'].c.computer_id==old['id'])).mappings().all()} if old else {}
                for line,item in group:
                    mac=_mac(item['MAC'])
                    adapter_type=ADAPTER_CN.get(_cell(item['网卡类型']))
                    if mac:
                        if not adapter_type:fail(f'第 {line} 行的网卡类型无效')
                        if mac not in seen_macs:adapters.append({'id':old_adapters.get(mac,{}).get('id'),'adapter_type':adapter_type,'mac_address':mac});seen_macs.add(mac)
                    elif _cell(item['关联 IP']):fail(f'第 {line} 行填写关联 IP 时必须填写 MAC')
                    if _cell(item['关联 IP']):assignments.append({'ip_address':_cell(item['关联 IP']),'adapter_mac':mac})
                    if _cell(item['配件类型']):
                        part=(_cell(item['配件类型']),_cell(item['配件备注']))
                        if part not in seen_parts:accessories.append({'accessory_type':part[0],'notes':part[1]});seen_parts.add(part)
                body={'asset_code':code,'device_type':_cell(first['设备类型']),'acceptance_month':first['验收日期'],'assignment_type':assignment_type,'site_id':site_id,
                    'user_id':user_id,'location':_cell(first['位置']),'operating_system':_cell(first['操作系统']),
                    'asset_status':ASSET_CN.get(_cell(first['状态']),_cell(first['状态']) or 'in_use'),'adapters':adapters,
                    'ip_assignments':assignments,'accessories':accessories,'notes':_cell(first['电脑备注'])}
                if old:body.update(version=old['version'],brand=old['brand'],model=old['model'],serial_number=old['serial_number'])
                result=save_computer(conn,user,body,old and old['id']);saved.append(result['id'])
            audit(conn,user,'computer.import','computer_assets',None,{'count':len(saved)})
            return out({'ok':True,'count':len(saved)})

    @app.get('/api/ipam-transfers/template')
    def ipam_template(request:Request):
        with engine.begin() as conn:require(conn,current_user(conn,request),'resource.import')
        return _download(_xlsx(IP_HEADERS),'static-ip-template.xlsx')

    @app.get('/api/ipam-transfers/export')
    def ipam_export(request:Request,network_key:str):
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.export');vlan,network,net=_managed_network(conn,user,network_key)
            require(conn,user,'resource.export',row(conn,'resources',vlan['id']));departments,users,_=_lookups(conn)
            records={v['ip_address']:dict(v) for v in conn.execute(sa.select(T['static_ip_records']).where(
                T['static_ip_records'].c.vlan_id==vlan['id'],T['static_ip_records'].c.network_id==network['network_id'])).mappings()}
            occ={v:k for k,v in OCCUPANCY_CN.items()};ver={v:k for k,v in VERIFICATION_CN.items()}
            rows=[];gateway=network.get('gateway')
            for address in net.hosts():
                ip_text=str(address)
                if ip_text==gateway:continue
                value=records.get(ip_text) or {'ip_address':ip_text,'occupancy_status':'available','verification_status':'pending'}
                active=value['occupancy_status']!='available'
                rows.append([vlan['site_name'],vlan['vlan_number'],network['cidr'],ip_text,
                    ('公共设备' if value.get('assignment_type')=='shared' else '个人使用') if active else '',
                    (departments.get(value.get('department_id')) or value.get('department_name_snapshot') or '') if active else '',
                    (users.get(value.get('user_id')) or value.get('user_name') or '') if active else '',
                    (value.get('device_type') or '') if active else '',(value.get('mac_address') or '') if active else '',
                    (value.get('location') or '') if active else '',value.get('notes') or '',
                    occ.get(value['occupancy_status'],value['occupancy_status']),ver.get(value['verification_status'],value['verification_status'])])
            audit(conn,user,'ipam.export','static_ip_records',None,{'network_key':network_key,'count':len(rows)})
        return _download(_xlsx(IP_HEADERS,rows),'static-ip-records.xlsx')

    @app.post('/api/ipam-transfers/import')
    async def ipam_import(request:Request,file:UploadFile=File(...)):
        data=_read_xlsx(await file.read(),IP_HEADERS)
        with engine.begin() as conn:
            user=current_user(conn,request);require(conn,user,'resource.import');departments,users,sites=_lookups(conn);options=_managed_options(conn,user);count=0
            people_review={'departed':0,'missing':0};department_review={'missing':0}
            adapters={v['mac_address']:dict(v) for v in conn.execute(sa.select(T['computer_network_adapters']).where(T['computer_network_adapters'].c.enabled==True)).mappings().all()}
            computers={v['id']:dict(v) for v in conn.execute(sa.select(T['computer_assets']).where(T['computer_assets'].c.is_archived==False)).mappings().all()};t=T['static_ip_records']
            for line,item in enumerate(data,2):
                site_id=_named_id(sites,item['区域'],'区域',True);vlan_number=_cell(item['VLAN ID'])
                if vlan_number.endswith('.0'):vlan_number=vlan_number[:-2]
                cidr=_cell(item['网段']);matches=[v for v in options if v['site_id']==site_id and str(v['vlan_number'])==vlan_number and v['cidr']==cidr]
                if len(matches)!=1:fail(f'第 {line} 行无法唯一匹配区域、VLAN ID 和网段')
                option=matches[0];vlan,network,net=_managed_network(conn,user,option['id']);require(conn,user,'resource.import',row(conn,'resources',vlan['id']))
                address=_cell(item['IP 地址'])
                try:ip=ipaddress.ip_address(address)
                except ValueError:fail(f'第 {line} 行 IP 地址无效')
                if ip not in net or ip in (net.network_address,net.broadcast_address) or address==network.get('gateway'):fail(f'第 {line} 行 IP 地址不在可登记范围')
                occupancy=OCCUPANCY_CN.get(_cell(item['状态']),_cell(item['状态']) or 'available');verification=VERIFICATION_CN.get(_cell(item['确认状态']),_cell(item['确认状态']) or 'pending')
                if occupancy not in OCCUPANCY or verification not in VERIFICATION:fail(f'第 {line} 行状态或确认状态无效')
                mac=_mac(item['MAC地址']);adapter=adapters.get(mac);computer=computers.get(adapter['computer_id']) if adapter else None
                if computer and computer['site_id']!=site_id:fail(f'第 {line} 行 MAC 所属电脑不在当前区域')
                responsible=occupancy in ('occupied','conflict');active=occupancy!='available';user_name=_cell(item['使用人'])
                assignment_type=(computer and computer.get('assignment_type')) or {'个人使用':'personal','公共设备':'shared'}.get(_cell(item['归属类型']),_cell(item['归属类型']) or 'personal')
                if assignment_type not in ('personal','shared'):fail(f'第 {line} 行归属类型无效')
                shared_computer=assignment_type=='shared'
                user_matches=[key for key,value in users.items() if value==user_name] if user_name else []
                user_id=user_matches[0] if len(user_matches)==1 else (computer and computer.get('user_id') if not user_name else None)
                if responsible and not user_id and not user_name and not shared_computer:fail(f'第 {line} 行已占用或冲突的 IP 必须填写使用人，或关联公共电脑资产')
                if shared_computer:user_id=None
                linked_user=row(conn,'users',user_id) if user_id else None
                person_status='missing' if active and not shared_computer and user_name and not user_id else 'departed' if active and not shared_computer and linked_user and linked_user['account_state']!='active' else ''
                department_name=_cell(item['部门']);department_matches=[key for key,value in departments.items() if value==department_name] if department_name else []
                if linked_user and linked_user['account_state']=='active' and department_name and departments.get(linked_user['department_id'])!=department_name:fail(f'第 {line} 行部门与系统人员资料不一致')
                department_id=(linked_user and linked_user['department_id']) or (department_matches[0] if len(department_matches)==1 else None)
                department_snapshot=department_name if shared_computer and department_name and not department_id else None
                department_status='missing' if active and shared_computer and department_snapshot else ''
                if person_status:
                    verification='pending';people_review[person_status]+=1
                if department_status:
                    verification='pending';department_review['missing']+=1
                old=conn.execute(sa.select(t).where(t.c.vlan_id==vlan['id'],t.c.network_id==network['network_id'],t.c.ip_address==address)).mappings().first();now=utcnow()
                device_type=((computer and computer['device_type']) or _cell(item['占用设备'])) if active else None
                location=(_cell(item['位置']) or (computer and computer['location'])) if active else None
                if responsible and assignment_type=='shared' and not device_type:fail(f'第 {line} 行公共设备必须填写占用设备')
                values={'occupancy_status':occupancy,'verification_status':verification,'assignment_type':assignment_type,'department_id':department_id if active else None,
                    'department_name_snapshot':department_snapshot if active else None,
                    'user_id':user_id if active else None,'user_name':(user_name or (linked_user and linked_user['display_name'])) if active and not shared_computer else None,'device_type':device_type,
                    'computer_id':computer['id'] if active and computer else None,'adapter_id':adapter['id'] if active and adapter else None,'mac_address':mac if active else None,
                    'location':location,'notes':_cell(item['备注']) or None,'verified_at':now if verification=='confirmed' else None,
                    'verified_by_user_id':user['id'] if verification=='confirmed' else None,'updated_at':now,'updated_by_user_id':user['id']}
                if old:conn.execute(t.update().where(t.c.id==old['id']).values(**values,version=old['version']+1));key=old['id'];action='import_updated'
                else:key=insert_row(conn,'static_ip_records',{'vlan_id':vlan['id'],'network_id':network['network_id'],'ip_address':address,**values},user['id'])['id'];action='import_created'
                saved=dict(row(conn,'static_ip_records',key));insert_row(conn,'static_ip_history',{'ip_record_id':key,'actor_user_id':user['id'],'action_code':action,'snapshot_json':json.dumps(jsonable(saved),ensure_ascii=False)});count+=1
            review={'people':people_review,'departments':department_review}
            audit(conn,user,'ipam.import','static_ip_records',None,{'count':count,'review':review});return out({'ok':True,'count':count,'people_review':people_review,'department_review':department_review})

    @app.post('/api/computers')
    def create_computer(request:Request,body:dict):
        with engine.begin() as conn:return out(save_computer(conn,current_user(conn,request),body),201)

    @app.put('/api/computers/{key}')
    def update_computer(request:Request,key:str,body:dict):
        with engine.begin() as conn:return out(save_computer(conn,current_user(conn,request),body,key))

    @app.delete('/api/computers/{key}')
    def delete_computer(request:Request,key:str,body:dict):
        with engine.begin() as conn:
            user=current_user(conn,request);old=row(conn,'computer_assets',key)
            if not old or old['is_archived']:fail('电脑设备不存在',404)
            require(conn,user,'resource.delete',old)
            if body.get('version')!=old['version']:fail('电脑设备资料已变化，请刷新',409)
            reason=_text(body.get('reason'),'删除原因',500,True)
            linked=conn.execute(sa.select(T['static_ip_records'].c.id).where(T['static_ip_records'].c.computer_id==key).limit(1)).first()
            if linked:fail('该电脑仍被静态 IP 记录关联，请先在静态 IP 管理中解除关联',409)
            conn.execute(T['computer_change_history'].delete().where(T['computer_change_history'].c.computer_id==key))
            conn.execute(T['computer_accessories'].delete().where(T['computer_accessories'].c.computer_id==key))
            conn.execute(T['computer_network_adapters'].delete().where(T['computer_network_adapters'].c.computer_id==key))
            conn.execute(T['computer_assets'].delete().where(T['computer_assets'].c.id==key))
            audit(conn,user,'computer.delete','computer_assets',key,reason=reason)
            return out({'ok':True})

    @app.get('/api/ipam/addresses')
    def addresses(request:Request,network_key:str,page:int=1,page_size:int=10,q:str='',occupancy_status:str='',verification_status:str='',asset_link:str='',consistency_status:str=''):
        if page_size not in (10,20,50):fail('每页条数仅支持 10、20、50')
        with engine.begin() as conn:
            user=current_user(conn,request);vlan,network,net=_managed_network(conn,user,network_key);page=max(1,page)
            t=T['static_ip_records'];records={v['ip_address']:dict(v) for v in conn.execute(sa.select(t).where(t.c.vlan_id==vlan['id'],t.c.network_id==network['network_id'])).mappings()}
            gateway=network.get('gateway');departments,users,_=_lookups(conn)
            user_states={v['id']:dict(v) for v in conn.execute(sa.select(T['users'].c.id,T['users'].c.account_state,T['users'].c.is_archived)).mappings()}
            department_states={v['id']:dict(v) for v in conn.execute(sa.select(T['departments'].c.id,T['departments'].c.name,T['departments'].c.is_archived)).mappings()}
            computers={v['id']:dict(v) for v in conn.execute(sa.select(T['computer_assets'])).mappings()}
            adapters={v['id']:dict(v) for v in conn.execute(sa.select(T['computer_network_adapters'])).mappings()}
            items=[];counts={k:0 for k in OCCUPANCY};pending=confirmed=needs_review=0
            for address in net.hosts():
                ip_text=str(address);record=records.get(ip_text)
                if record:item=record
                elif ip_text==gateway:item={'id':None,'version':0,'ip_address':ip_text,'occupancy_status':'reserved','verification_status':'confirmed','notes':'VLAN 网关','system_reserved':True}
                else:item={'id':None,'version':0,'ip_address':ip_text,'occupancy_status':'available','verification_status':'pending'}
                counts[item['occupancy_status']]+=1
                if item['occupancy_status']!='available':
                    if item['verification_status']=='confirmed':confirmed+=1
                    else:pending+=1
                linked_department=department_states.get(item.get('department_id'))
                display_department=departments.get(item.get('department_id')) or item.get('department_name_snapshot') or (linked_department and linked_department['name'])
                item={**item,'department_name':display_department,'display_user':users.get(item.get('user_id')) or item.get('user_name')}
                linked_person=user_states.get(item.get('user_id'))
                person_missing=(bool(item.get('user_name')) and not item.get('user_id')) or (bool(item.get('user_id')) and (not linked_person or linked_person['is_archived']))
                item['person_status']='missing' if person_missing else 'departed' if linked_person and linked_person['account_state']!='active' else ''
                item['department_status']='missing' if item.get('assignment_type')=='shared' and display_department and (not linked_department or linked_department['is_archived']) else ''
                computer=computers.get(item.get('computer_id'));item['computer_name']=computer and computer['asset_code'];issues=[]
                if item['person_status']:issues.append('人员不存在' if item['person_status']=='missing' else '人员已离职')
                if item['department_status']:issues.append('部门不存在')
                if computer and item['occupancy_status']!='available':
                    adapter=adapters.get(item.get('adapter_id'));computer_user=users.get(computer.get('user_id')) or computer.get('user_name')
                    if item.get('assignment_type')!='shared' and item.get('department_id')!=computer.get('department_id'):issues.append('部门不一致')
                    if item.get('assignment_type')!='shared' and (item.get('display_user') or '')!=(computer_user or ''):issues.append('使用人不一致')
                    if item.get('assignment_type')!=computer.get('assignment_type'):issues.append('归属类型不一致')
                    if (item.get('location') or '')!=(computer.get('location') or ''):issues.append('位置不一致')
                    if item.get('device_type')!=computer.get('device_type'):issues.append('设备类型不一致')
                    if not adapter or adapter.get('computer_id')!=computer['id']:issues.append('未关联电脑网卡')
                    elif item.get('mac_address')!=adapter.get('mac_address'):issues.append('MAC 与所选网卡不一致')
                    item['asset_link_status']='linked';item['consistency_status']='conflict' if issues else 'consistent'
                elif item['occupancy_status']!='available' and not item.get('system_reserved'):
                    candidate=_mac_asset_candidate(conn,item,vlan['site_id'])
                    if candidate:
                        candidate_computer=computers.get(candidate['computer_id'])
                        item.update(candidate_computer_id=candidate['computer_id'],candidate_adapter_id=candidate['id'],
                            candidate_asset_code=candidate['asset_code'],candidate_device_type=candidate['device_type'],
                            candidate_user=users.get(candidate_computer and candidate_computer.get('user_id')) or '公共设备')
                        item['asset_link_status']='candidate';item['consistency_status']='pending'
                        issues.append('MAC 已匹配电脑资产 '+candidate['asset_code']+'，待确认关联')
                    elif item.get('device_type') in COMPUTER_DEVICE_TYPES or item.get('device_type')=='电脑':
                        item['asset_link_status']='unlinked';item['consistency_status']='pending';issues.append('电脑类设备尚未关联电脑资产')
                    elif item.get('device_type'):
                        item['asset_link_status']='non_computer';item['consistency_status']='consistent'
                    else:
                        item['asset_link_status']='unlinked';item['consistency_status']='pending';issues.append('尚未确认占用设备类型')
                else:item['asset_link_status']='';item['consistency_status']=''
                unresolved={'人员不存在','人员已离职','部门不存在'}
                if (item['person_status'] or item['department_status']) and item['consistency_status'] in ('','consistent','conflict') and issues and all(issue in unresolved for issue in issues):item['consistency_status']='pending'
                item['consistency_issues']=issues
                if item['consistency_status'] in ('pending','conflict'):needs_review+=1
                haystack=json.dumps(jsonable(item),ensure_ascii=False).casefold()
                if q and q.casefold() not in haystack:continue
                if occupancy_status and item['occupancy_status']!=occupancy_status:continue
                if verification_status and item['verification_status']!=verification_status:continue
                if asset_link and not (asset_link=='unlinked' and item['asset_link_status'] in ('unlinked','candidate')) and item['asset_link_status']!=asset_link:continue
                if consistency_status and item['consistency_status']!=consistency_status:continue
                items.append(item)
            total=len(items);start=(page-1)*page_size
            return out({'items':items[start:start+page_size],'total':total,'page':page,'page_size':page_size,
                'network':{'id':network_key,'site_name':vlan['site_name'],'vlan_number':vlan['vlan_number'],**network},
                'stats':{'total':sum(counts.values()),**counts,'pending':pending,'confirmed':confirmed,'needs_review':needs_review}})

    @app.post('/api/ipam/records')
    def save_ip(request:Request,body:dict):
        with engine.begin() as conn:
            user=current_user(conn,request);vlan,network,net=_managed_network(conn,user,str(body.get('network_key') or ''))
            address=str(body.get('ip_address') or '')
            try:ip=ipaddress.ip_address(address)
            except ValueError:fail('IP 地址无效')
            if ip not in net or ip in (net.network_address,net.broadcast_address):fail('IP 地址不属于可用主机范围')
            if address==network.get('gateway'):fail('网关地址由系统保留，不能修改')
            occupancy=body.get('occupancy_status') or 'occupied';verification=body.get('verification_status') or 'pending'
            if occupancy not in OCCUPANCY or verification not in VERIFICATION:fail('IP状态无效')
            t=T['static_ip_records'];old=conn.execute(sa.select(t).where(t.c.vlan_id==vlan['id'],t.c.network_id==network['network_id'],t.c.ip_address==address)).mappings().first()
            require(conn,user,'resource.edit',row(conn,'resources',vlan['id']))
            if old and body.get('version')!=old['version']:fail('IP资料已变化，请刷新',409)
            computer=row(conn,'computer_assets',body.get('computer_id')) if body.get('computer_id') else None
            adapter=row(conn,'computer_network_adapters',body.get('adapter_id')) if body.get('adapter_id') else None
            if computer and (computer['is_archived'] or computer['site_id']!=vlan['site_id']):fail('关联电脑须属于当前区域')
            if adapter and (not computer or adapter['computer_id']!=computer['id']):fail('所选网卡不属于关联电脑')
            preserve_on_link=bool(old and old['occupancy_status']!='available' and computer and not old.get('computer_id'))
            assignment_type=(computer and computer.get('assignment_type')) or body.get('assignment_type') or (old and old.get('assignment_type')) or 'personal'
            if assignment_type not in ('personal','shared'):fail('IP 归属类型无效')
            user_id=body.get('user_id') or (old.get('user_id') if preserve_on_link else None) or (computer and computer['user_id'])
            if assignment_type=='shared':user_id=None
            linked_user=row(conn,'users',user_id) if user_id else None
            if occupancy in ('occupied','conflict') and (not linked_user or linked_user['is_archived']) and assignment_type!='shared':fail('个人使用设备占用 IP 时必须选择系统人员')
            requested_department_id=body.get('department_id') if assignment_type=='shared' else None
            selected_department=row(conn,'departments',requested_department_id) if requested_department_id else None
            if requested_department_id and (not selected_department or selected_department['is_archived']):fail('请选择有效的管理/使用部门')
            department_id=selected_department['id'] if assignment_type=='shared' and selected_department else linked_user['department_id'] if linked_user else None
            snapshot_source=body.get('department_name_snapshot') if 'department_name_snapshot' in body else old and old.get('department_name_snapshot')
            department_snapshot=_text(snapshot_source,'原始部门名称',128) if assignment_type=='shared' and not department_id else None
            if department_snapshot:verification='pending'
            active=occupancy!='available'
            submitted_mac=_mac(body.get('mac_address'))
            if active and computer:
                if not adapter:fail('关联电脑时必须选择该电脑的具体网卡')
                if submitted_mac and submitted_mac!=adapter['mac_address']:fail('IP 地址表中的 MAC 与所选电脑网卡不一致，不能建立关联')
                if old and old['occupancy_status']!='available' and old.get('computer_id')!=computer['id'] and old.get('mac_address')!=adapter['mac_address']:
                    fail('该 IP 原记录的 MAC 与所选电脑网卡不一致，请先核实 MAC 地址',409)
            device_type=((old.get('device_type') if preserve_on_link else None) or (computer and computer['device_type']) or _text(body.get('device_type'),'适用设备',64)) if active else None
            location=(_text(body.get('location'),'位置',256) or (old.get('location') if preserve_on_link else None) or (computer and computer['location'])) if active else None
            if occupancy in ('occupied','conflict') and assignment_type=='shared' and not device_type:fail('公共设备占用 IP 时必须填写占用设备类型')
            values={'occupancy_status':occupancy,'verification_status':verification,'assignment_type':assignment_type,
                'department_id':department_id if active else None,'user_id':user_id if active else None,
                'department_name_snapshot':department_snapshot if active else None,
                'user_name':None,
                'device_type':device_type,
                'computer_id':computer['id'] if active and computer else None,'adapter_id':adapter['id'] if active and adapter else None,
                'mac_address':((adapter and adapter['mac_address']) or submitted_mac) if active else None,
                'location':location,
                'notes':_text(body.get('notes'),'备注',10000),
                'verified_at':utcnow() if verification=='confirmed' else None,'verified_by_user_id':user['id'] if verification=='confirmed' else None}
            now=utcnow()
            if old:
                conn.execute(t.update().where(t.c.id==old['id'],t.c.version==old['version']).values(**values,version=old['version']+1,updated_at=now,updated_by_user_id=user['id']));key=old['id'];action='updated'
            else:
                key=insert_row(conn,'static_ip_records',{'vlan_id':vlan['id'],'network_id':network['network_id'],'ip_address':address,**values},user['id'])['id'];action='created'
            saved=dict(row(conn,'static_ip_records',key));insert_row(conn,'static_ip_history',{'ip_record_id':key,'actor_user_id':user['id'],'action_code':action,'snapshot_json':json.dumps(jsonable(saved),ensure_ascii=False)})
            audit(conn,user,'ipam.save','static_ip_records',key,{'ip_address':address,'status':occupancy})
            return out(saved)

    @app.post('/api/ipam/records/{key}/reconcile')
    def reconcile_ip(request:Request,key:str,body:dict):
        with engine.begin() as conn:
            user=current_user(conn,request);record=row(conn,'static_ip_records',key)
            if not record:fail('IP记录不存在',404)
            vlan=row(conn,'resources',record['vlan_id']);require(conn,user,'resource.edit',vlan)
            if body.get('version')!=record['version']:fail('IP资料已变化，请刷新',409)
            computer=row(conn,'computer_assets',record['computer_id']) if record.get('computer_id') else None
            adapter=row(conn,'computer_network_adapters',record['adapter_id']) if record.get('adapter_id') else None
            if not computer and not adapter:
                candidate=_mac_asset_candidate(conn,record,vlan['site_id'])
                if candidate:
                    computer=row(conn,'computer_assets',candidate['computer_id']);adapter=row(conn,'computer_network_adapters',candidate['id'])
            if not computer or computer['is_archived']:fail('该 IP 尚未关联有效电脑资产')
            if not adapter or adapter['computer_id']!=computer['id']:fail('请先为该 IP 重新选择电脑网卡')
            action=body.get('action');now=utcnow();ip_table=T['static_ip_records']
            if action=='from_computer':
                shared=computer.get('assignment_type')=='shared'
                values={'assignment_type':computer.get('assignment_type') or 'personal',
                    'department_id':record.get('department_id') if shared else computer['department_id'],
                    'department_name_snapshot':record.get('department_name_snapshot') if shared else None,
                    'user_id':computer['user_id'],'user_name':computer['user_name'],
                    'device_type':computer['device_type'],'computer_id':computer['id'],'adapter_id':adapter['id'],
                    'mac_address':adapter['mac_address'],'location':computer['location'],
                    'version':record['version']+1,'updated_at':now,'updated_by_user_id':user['id']}
                conn.execute(ip_table.update().where(ip_table.c.id==key).values(**values));updated=[key]
            elif action=='to_computer':
                require(conn,user,'resource.edit',computer);before=_computer(conn,user,computer['id']);computer_table=T['computer_assets']
                assignment_type=record.get('assignment_type') or 'personal'
                if assignment_type not in ('personal','shared'):fail('IP 记录的归属类型无效')
                target_user=row(conn,'users',record['user_id']) if record.get('user_id') else None
                if assignment_type=='personal' and (not target_user or target_user['is_archived']):fail('IP 记录没有有效使用人，不能写入电脑台账')
                if assignment_type=='shared':target_user=None
                target_department=target_user['department_id'] if target_user else None
                target_device_type=record.get('device_type')
                if target_device_type not in COMPUTER_DEVICE_TYPES:fail('IP 记录的设备类型必须是笔记本或台式机，才能写入电脑台账')
                conn.execute(computer_table.update().where(computer_table.c.id==computer['id'],computer_table.c.version==computer['version']).values(
                    assignment_type=assignment_type,department_id=target_department,user_id=target_user and target_user['id'],user_name=None,
                    device_type=target_device_type,location=record['location'],
                    version=computer['version']+1,updated_at=now,updated_by_user_id=user['id']))
                conn.execute(ip_table.update().where(ip_table.c.id==key).values(computer_id=computer['id'],adapter_id=adapter['id']))
                linked=conn.execute(sa.select(ip_table).where(ip_table.c.computer_id==computer['id'],ip_table.c.occupancy_status!='available')).mappings().all();updated=[]
                for linked_record in linked:
                    linked_adapter=row(conn,'computer_network_adapters',linked_record['adapter_id']) if linked_record.get('adapter_id') else None
                    record_department=linked_record.get('department_id') if assignment_type=='shared' else target_department
                    record_department_snapshot=linked_record.get('department_name_snapshot') if assignment_type=='shared' else None
                    conn.execute(ip_table.update().where(ip_table.c.id==linked_record['id']).values(assignment_type=assignment_type,
                        department_id=record_department,department_name_snapshot=record_department_snapshot,user_id=target_user and target_user['id'],user_name=None,
                        device_type=target_device_type,location=record['location'],mac_address=linked_adapter and linked_adapter['mac_address'],
                        version=linked_record['version']+1,updated_at=now,updated_by_user_id=user['id']));updated.append(linked_record['id'])
                after=_computer(conn,user,computer['id']);summary=f"通过 IP {record['ip_address']} 核对电脑资料：设备类型 {before.get('device_type') or '未填写'} → {after.get('device_type') or '未填写'}；部门 {before.get('department_name') or '未填写'} → {after.get('department_name') or '未填写'}；使用人 {before.get('display_user') or '未分配'} → {after.get('display_user') or '未分配'}；位置 {before.get('location') or '未填写'} → {after.get('location') or '未填写'}"
                _record_computer_history(conn,user,computer['id'],'ip_reconciled',summary)
            else:fail('核对处理方式无效')
            for record_id in updated:
                saved=dict(row(conn,'static_ip_records',record_id));insert_row(conn,'static_ip_history',{'ip_record_id':record_id,'actor_user_id':user['id'],'action_code':'reconciled','snapshot_json':json.dumps(jsonable(saved),ensure_ascii=False)})
            audit(conn,user,'ipam.reconcile','static_ip_records',key,{'action':action,'computer_id':computer['id']})
            return out({'ok':True,'updated':len(updated)})

    @app.get('/api/ipam/records/{key}/history')
    def ip_history(request:Request,key:str):
        with engine.begin() as conn:
            user=current_user(conn,request);record=row(conn,'static_ip_records',key)
            if not record:fail('IP记录不存在',404)
            vlan=row(conn,'resources',record['vlan_id']);require(conn,user,'resource.view',vlan)
            rows=conn.execute(sa.select(T['static_ip_history']).where(T['static_ip_history'].c.ip_record_id==key).order_by(T['static_ip_history'].c.created_at.desc())).mappings()
            return out({'items':[dict(v) for v in rows]})
