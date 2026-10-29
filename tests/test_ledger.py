from io import BytesIO
from datetime import date
import unittest
from openpyxl import load_workbook
from tests import test_business as business
from portal.database import TABLES as T
import sqlalchemy as sa


class LedgerTests(unittest.TestCase):
    setUp=business.BusinessTests.setUp
    login=business.BusinessTests.login
    add_user=business.BusinessTests.add_user

    def save(self,kind,body,key=None,status=201):
        r=self.client.request('PUT' if key else 'POST','/api/ledger/'+kind+('/'+key if key else ''),json=body)
        self.assertEqual(r.status_code,status,r.text);return r.json()

    def base(self,kind,**values):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        return dict(site_id=site,name='测试'+kind,**values)

    def test_vlan_wifi_live_reference_impact_and_delete(self):
        v=self.save('vlan',self.base('VLAN',vlan_number=100,networks=[{'cidr':'10.20.10.0/24','gateway':'10.20.10.1'}]))
        w=self.save('wifi',self.base('WiFi',ssid='Company-WiFi',vlan_ids=[v['id']]))
        self.assertEqual(w['vlan_details'][0]['networks'][0]['gateway'],'10.20.10.1')
        update=dict(version=v['version'],networks=[{'cidr':'10.20.20.0/24','gateway':'10.20.20.1'}])
        self.save('vlan',update,v['id'],409)
        changed=self.save('vlan',{**update,'acknowledge_impacts':True},v['id'],200)
        w=self.client.get('/api/ledger/item/'+w['id']).json()
        self.assertEqual(w['vlan_details'][0]['networks'][0]['cidr'],'10.20.20.0/24')
        r=self.client.request('DELETE','/api/ledger/item/'+v['id'],json={'version':changed['version'],'confirm':v['id'],'reason':'引用保护测试'})
        self.assertEqual(r.status_code,409,r.text);self.assertIn('WiFi',r.text)
        self.save('vlan',{'version':v['version'],'name':'旧版本更新'},v['id'],409)

    def test_circuit_gift_many_ips_aws_two_tunnels_and_device(self):
        v=self.save('vlan',self.base('VLAN',vlan_number=101,networks=[{'cidr':'10.20.10.0/24'}]))
        d=self.save('device',self.base('设备',model='S5735',manufacturer='华为',esn='ESN-01',management_ip='10.20.10.2',warranty_start='2026-01-01',warranty_end='2027-01-01'))
        self.save('device',self.base('错误设备',management_ip='not-an-ip'),status=422)
        lines=[{'line_type':'primary','number':'C-01','downstream':'100','upstream':'100','usable_ips':'203.0.113.2\n203.0.113.3','mask':'29','gateway':'203.0.113.1','interface':'GE0/0/1'},
               {'line_type':'gift','number':'G-01','usable_ips':'198.51.100.2','mask':'255.255.255.248','gateway':'198.51.100.1'}]
        c=self.save('circuit',self.base('线路',vendor_name='测试运营商',lines=lines))
        self.assertEqual([v['ip_address'] for v in c['connections'][:2]],['203.0.113.2','203.0.113.3'])
        a=self.save('vpn',self.base('AWS',circuit_id=c['id'],local_vlan_ids=[v['id']],local_gateway='203.0.113.2',remote_networks=[{'cidr':'172.20.0.0/16','notes':'AWS业务网段'}],tunnels=[{'remote_gateway':'192.0.2.1','registered_state':'up'},{'remote_gateway':'192.0.2.2','registered_state':'disabled'}]))
        self.assertNotIn('circuit_id',a)
        self.assertEqual(a['vlan_details'][0]['vlan_number'],101);self.assertEqual(a['tunnels'][1]['registered_state'],'disabled')
        self.save('wifi',self.base('WiFi',ssid='Roundtrip-WiFi',vlan_ids=[v['id']]))
        for kind in ('wifi','device','circuit','vpn'):
            raw=self.client.get('/api/ledger/'+kind+'/export').content
            r=self.client.post('/api/ledger/'+kind+'/imports',files={'file':(kind+'.xlsx',raw)},data={'mode':'upsert'})
            self.assertEqual(r.status_code,201,r.text);batch=r.json()
            staged=self.client.get('/api/imports/batches/'+batch['id']).json()
            self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
            self.assertTrue(all(v['source_row_key'].isdigit() for v in staged['rows']))
            r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)

    def test_ledger_template_atomic_import_export_and_conflict(self):
        raw=self.client.get('/api/ledger/vlan/template').content;wb=load_workbook(BytesIO(raw));headers=[c.value for c in wb['数据'][1]]
        values={'区域':'外仓','VLAN ID':200,'网段 CIDR':'10.50.0.0/24','网关':'10.50.0.1','启用':'是'}
        wb['数据'].append([values.get(k) for k in headers]);out=BytesIO();wb.save(out)
        r=self.client.post('/api/ledger/vlan/imports',files={'file':('vlan.xlsx',out.getvalue())});self.assertEqual(r.status_code,201,r.text);batch=r.json()
        self.assertEqual(self.client.get('/api/ledger/vlan').json()['total'],0)
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        exported=self.client.get('/api/ledger/vlan/export');self.assertEqual(exported.status_code,200)
        r=self.client.post('/api/ledger/vlan/imports',files={'file':('vlan.xlsx',exported.content)},data={'mode':'upsert'});self.assertEqual(r.status_code,201,r.text)
        staged=self.client.get('/api/imports/batches/'+r.json()['id']).json();self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
        old=self.client.get('/api/ledger/vlan').json()['items'][0]
        self.save('vlan',{'version':old['version'],'name':'并发修改'},old['id'],200)
        r=self.client.post('/api/imports/batches/'+r.json()['id']+'/commit',json={});self.assertEqual(r.status_code,409,r.text)

    def test_column_templates_group_multiple_rows_and_reject_conflicting_shared_values(self):
        import json
        raw=self.client.get('/api/ledger/vlan/template').content
        wb=load_workbook(BytesIO(raw));sheet=wb['数据'];headers=[c.value for c in sheet[1]]
        self.assertNotIn('网段与网关',headers)
        for cidr,purpose in [('10.111.0.0/24','办公 | A\n二楼'),('10.112.0.0/24','生产')]:
            values={'区域':'外仓','VLAN ID':401,'网段 CIDR':cidr,'网段用途':purpose,'启用':'是'}
            sheet.append([values.get(k) for k in headers])
        out=BytesIO();wb.save(out)
        batch=self.client.post('/api/ledger/vlan/imports',files={'file':('new.xlsx',out.getvalue())}).json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json']),{'module':'resources','ledger_kind':'vlan','rows':1,'errors':0})
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,200,result.text)
        raw=self.client.get('/api/ledger/vlan/export').content;wb=load_workbook(BytesIO(raw));sheet=wb['数据']
        self.assertEqual(sheet.max_row,3)
        self.assertEqual(sheet.cell(2,headers.index('网段用途')+1).value,'办公 | A\n二楼')
        self.assertEqual(sheet.cell(2,1).value,sheet.cell(3,1).value)
        sheet.cell(3,headers.index('备注')+1,'不同备注')
        sheet.cell(2,headers.index('备注')+1,'首行备注')
        out=BytesIO();wb.save(out)
        batch=self.client.post('/api/ledger/vlan/imports',files={'file':('conflict.xlsx',out.getvalue())},data={'mode':'upsert'}).json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],1)
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,409,result.text)

    def test_new_wifi_grouped_rows_select_separate_vlan_columns(self):
        import json
        v=self.save('vlan',self.base('模板网段',vlan_number=402,networks=[{'cidr':'10.121.0.0/24'},{'cidr':'10.122.0.0/24'}]))
        wb=load_workbook(BytesIO(self.client.get('/api/ledger/wifi/template').content));sheet=wb['数据'];headers=[c.value for c in sheet[1]]
        self.assertEqual(headers[:10],['区域','VLAN ID','网段','SSID','密码','认证方式','用途','覆盖位置','启用','备注'])
        self.assertTrue(sheet.column_dimensions['K'].hidden);self.assertTrue(sheet.column_dimensions['L'].hidden)
        for index,n in enumerate(v['networks']):
            values={'导入分组（系统辅助列）':'office','VLAN ID':402,'网段':n['cidr']}
            if index==0:values.update({'区域':'外仓','SSID':'GroupedWiFi','认证方式':'不认证','覆盖位置':'办公楼 1-3 层'})
            sheet.append([values.get(k) for k in headers])
        out=BytesIO();wb.save(out)
        batch=self.client.post('/api/ledger/wifi/imports',files={'file':('wifi.xlsx',out.getvalue())}).json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,200,result.text)
        w=self.client.get('/api/ledger/wifi').json()['items'][0];self.assertEqual(len(w['vlan_network_ids']),2)
        self.assertEqual(w['coverage'],'办公楼 1-3 层')

    def test_new_vpn_grouped_rows_keep_multiple_vlan_networks_in_one_record(self):
        import json
        site=self.client.get('/api/lookups').json()['sites'][0]
        first=self.save('vlan',{'site_id':site['id'],'vlan_number':451,'networks':[{'cidr':'10.151.0.0/24','purpose':'办公'}]})
        second=self.save('vlan',{'site_id':site['id'],'vlan_number':452,'networks':[{'cidr':'10.152.0.0/24','purpose':'生产'}]})
        response=self.client.get('/api/ledger/vpn/template');self.assertEqual(response.status_code,200,response.text)
        wb=load_workbook(BytesIO(response.content));sheet=wb['数据'];headers=[c.value for c in sheet[1]]
        for index,vlan in enumerate((first,second)):
            values={'导入分组（同一条资料填写相同值）':'VPN001','本地 VLAN 网段·网络范围':'default',
                '本地 VLAN 网段·VLAN ID':vlan['vlan_number'],'本地 VLAN 网段·网段 CIDR':vlan['networks'][0]['cidr']}
            if index==0:values.update({'区域':site['name'],'启用':'是','隧道1·对端网关':'192.0.2.10','隧道1·登记状态':'正常'})
            sheet.append([values.get(name) for name in headers])
        output=BytesIO();wb.save(output)
        preview=self.client.post('/api/ledger/vpn/imports',files={'file':('vpn.xlsx',output.getvalue())})
        self.assertEqual(preview.status_code,201,preview.text);batch=preview.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json']),{'module':'resources','ledger_kind':'vpn','rows':1,'errors':0})
        committed=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(committed.status_code,200,committed.text)
        vpn=self.client.get('/api/ledger/vpn').json()['items'][0]
        self.assertEqual(len(vpn['local_vlan_network_ids']),2);self.assertEqual(len(vpn['local_network_details']),2)

    def test_vpn_any_network_mode_clears_vlan_links_and_imports(self):
        import json
        site=self.client.get('/api/lookups').json()['sites'][0]
        vlan=self.save('vlan',{'site_id':site['id'],'vlan_number':453,'networks':[{'cidr':'10.153.0.0/24','purpose':'测试网段'}]})
        selected=vlan['id']+':'+vlan['networks'][0]['network_id']
        vpn=self.save('vpn',{'site_id':site['id'],'remote_site_id':site['id'],
            'local_network_mode':'any','local_vlan_network_ids':[selected],
            'remote_network_mode':'any','remote_vlan_network_ids':[selected],
            'remote_networks':[{'cidr':'172.31.0.0/16','notes':'应被 Any 清除'}]})
        self.assertEqual(vpn['local_network_mode'],'any');self.assertEqual(vpn['remote_network_mode'],'any')
        self.assertEqual(vpn['local_vlan_network_ids'],[]);self.assertEqual(vpn['remote_vlan_network_ids'],[])
        self.assertEqual(vpn['remote_networks'],[])
        self.assertEqual(vpn['local_network_details'],[{'cidr':'0.0.0.0/0','purpose':'Any','any':True}])
        self.assertEqual(vpn['remote_cidrs'],'0.0.0.0/0')

        wb=load_workbook(BytesIO(self.client.get('/api/ledger/vpn/template').content));sheet=wb['数据'];headers=[c.value for c in sheet[1]]
        self.assertNotIn('承载线路',headers)
        values={'区域':site['name'],'本地 VPN 网段':'Any','对端 VPN 网段':'0.0.0.0/0','启用':'是'}
        sheet.append([values.get(name) for name in headers]);output=BytesIO();wb.save(output)
        preview=self.client.post('/api/ledger/vpn/imports',files={'file':('vpn-any.xlsx',output.getvalue())})
        self.assertEqual(preview.status_code,201,preview.text);batch=preview.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
        committed=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(committed.status_code,200,committed.text)
        imported=[v for v in self.client.get('/api/ledger/vpn').json()['items'] if v['id']!=vpn['id']][0]
        self.assertEqual(imported['local_network_details'][0]['cidr'],'0.0.0.0/0')
        self.assertEqual(imported['remote_network_details'][0]['cidr'],'0.0.0.0/0')

    def device_import(self,items,mode='create_only'):
        wb=load_workbook(BytesIO(self.client.get('/api/ledger/device/template').content))
        sheet=wb['数据'];headers=[c.value for c in sheet[1]]
        for item in items:sheet.append([item.get(k) for k in headers])
        out=BytesIO();wb.save(out)
        result=self.client.post('/api/ledger/device/imports',files={'file':('device.xlsx',out.getvalue())},data={'mode':mode})
        self.assertEqual(result.status_code,201,result.text)
        return result.json()

    def test_device_ip_uplinks_same_batch_reverse_order_and_export(self):
        import json
        batch=self.device_import([
            {'区域':'外仓','设备型号':'access','管理 IP':'10.150.0.3','上联设备':' distribution · 10.150.0.2 '},
            {'区域':'外仓','设备型号':'distribution','管理 IP':'10.150.0.2','上联设备':'10.150.0.1'},
            {'区域':'外仓','设备型号':'core','管理 IP':'10.150.0.1'},
        ])
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
        self.assertEqual(self.client.get('/api/ledger/device').json()['total'],0,'preview must not save devices')
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,200,result.text)
        items={v['management_ip']:v for v in self.client.get('/api/ledger/device').json()['items']}
        self.assertEqual(items['10.150.0.3']['uplink_device_id'],items['10.150.0.2']['id'])
        self.assertEqual(items['10.150.0.2']['uplink_device_id'],items['10.150.0.1']['id'])
        wb=load_workbook(BytesIO(self.client.get('/api/ledger/device/export').content))
        headers=[c.value for c in wb['数据'][1]];index=headers.index('上联设备')
        self.assertEqual({r[index] for r in list(wb['数据'].values)[1:] if r[index]},{'10.150.0.1','10.150.0.2'})

    def test_device_ip_uplinks_duplicates_wrong_area_and_cycles_blocked(self):
        import json
        for label in ('one','two'):self.save('device',self.base(label,management_ip='10.151.0.1'))
        for items in [
            [{'区域':'外仓','管理 IP':'10.151.0.2','上联设备':'10.151.0.1'}],
            [{'区域':'外仓','管理 IP':'10.152.0.1','上联设备':'10.152.0.2'}, {'区域':'外仓','管理 IP':'10.152.0.2','上联设备':'10.152.0.1'}],
        ]:
            batch=self.device_import(items);staged=self.client.get('/api/imports/batches/'+batch['id']).json()
            self.assertGreater(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
            result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,409,result.text)
        other=self.client.post('/api/data/sites',json={'name':'异地区域'});self.assertEqual(other.status_code,201,other.text)
        self.save('device',{'site_id':other.json()['id'],'management_ip':'10.153.0.1'})
        batch=self.device_import([{'区域':'外仓','管理 IP':'10.153.0.2','上联设备':'10.153.0.1'}])
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],1,staged)

    def test_device_ip_upsert_preserves_identity_password_and_prevents_duplicates(self):
        import json
        core=self.save('device',self.base('core',management_ip='10.154.0.1',login_password='Existing device secret'))
        items=[{'区域':'外仓','管理 IP':'10.154.0.2','上联设备':'10.154.0.1'},
               {'区域':'外仓','管理 IP':'10.154.0.1','设备型号':'updated core'}]
        batch=self.device_import(items)
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertGreater(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
        self.assertEqual(self.client.get('/api/ledger/device').json()['total'],1)
        batch=self.device_import(items,mode='upsert')
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(json.loads(staged['batch']['summary_json'])['errors'],0,staged)
        self.assertEqual(self.client.get('/api/ledger/item/'+core['id']).json()['model'],None)
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,200,result.text)
        devices={v['management_ip']:v for v in self.client.get('/api/ledger/device').json()['items']}
        self.assertEqual(len(devices),2)
        self.assertEqual(devices['10.154.0.1']['id'],core['id'])
        self.assertEqual(devices['10.154.0.2']['uplink_device_id'],core['id'])
        actual=self.client.get('/api/ledger/item/'+core['id']).json()
        self.assertEqual(actual['model'],'updated core');self.assertEqual(actual['login_password'],'Existing device secret')

    def test_legacy_device_preserved_and_employee_access_denied(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        r=self.client.post('/api/resources/device',json={'resource':{'name':'原有设备','site_id':site},'details':{'device_type_code':'switch','model':'Old Model'}})
        self.assertEqual(r.status_code,201,r.text);key=r.json()['resource']['id']
        old=self.client.get('/api/ledger/item/'+key).json();self.assertEqual(old['model'],'Old Model')
        changed=self.save('device',{'version':old['version'],'esn':'NEW-ESN'},key,200)
        self.assertEqual(changed['id'],key);self.assertEqual(changed['model'],'Old Model')
        self.login('employee-a')
        for path in ('/api/ledger/schema','/api/ledger/options','/api/ledger/device','/api/ledger/item/'+key):self.assertEqual(self.client.get(path).status_code,403)
        self.assertEqual(self.client.post('/api/ledger/device',json=self.base('越权设备')).status_code,403)

    def test_export_scope_is_independent_of_view_scope(self):
        for label,owner in [('可导出设备',self.it_a),('仅可查看设备',self.it_b)]:
            obj=self.save('device',self.base(label))
            with self.app.state.engine.begin() as conn:conn.execute(T['resources'].update().where(T['resources'].c.id==obj['id']).values(owner_user_id=owner))
        role=self.client.post('/api/roles',json={'code':'ledger_scope','name':'台账范围测试','permissions':[
            {'code':'resource.view','scope':'all'},{'code':'resource.export','scope':'assigned'}]})
        self.assertEqual(role.status_code,201,role.text)
        user=next(v for v in self.client.get('/api/users').json()['items'] if v['id']==self.it_a)
        r=self.client.put('/api/users/'+self.it_a,json={'version':user['version'],'role_ids':[role.json()['id']]})
        self.assertEqual(r.status_code,200,r.text);self.login('it-a')
        self.assertEqual(self.client.get('/api/ledger/device').json()['total'],2)
        r=self.client.get('/api/ledger/device/export');self.assertEqual(r.status_code,200,r.text)
        wb=load_workbook(BytesIO(r.content));self.assertEqual(wb['数据'].max_row,2)

    def test_simplified_forms_multiple_networks_and_wifi_password(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        v=self.save('vlan',{'site_id':site,'vlan_number':320,'networks':[{'cidr':'10.32.0.0/24','gateway':'10.32.0.1'},{'cidr':'10.33.0.0/24','gateway':'10.33.0.1'}]})
        self.assertTrue(v['name']);self.assertEqual(len(v['networks']),2)
        schemas=self.client.get('/api/ledger/schema').json()['modules']
        for spec in schemas:self.assertTrue(next(f for f in spec['fields'] if f['name']=='name')['ui_hidden'])
        vlan=next(s for s in schemas if s['kind']=='vlan')
        self.assertEqual([f['name'] for f in next(f for f in vlan['fields'] if f['name']=='networks')['columns']],['cidr','gateway','purpose','ipam_enabled','network_id'])
        password='QA WiFi test password! '
        w=self.save('wifi',{'site_id':site,'ssid':'Test-New-WiFi','auth_method':'password_mac','wifi_password':password,'vlan_ids':[v['id']]})
        self.assertEqual(w['wifi_password'],password)
        with self.app.state.engine.connect() as conn:
            stored=conn.execute(sa.select(T['ledger_profiles'].c.data_json).where(T['ledger_profiles'].c.resource_id==w['id'])).scalar_one()
            self.assertNotIn(password,stored);self.assertIn('wifi_password_ciphertext',stored)
        for path in ('/api/ledger/wifi','/api/ledger/options'):
            self.assertNotIn(password,self.client.get(path).text)
        read=self.client.get('/api/ledger/item/'+w['id']).json()
        self.assertEqual(read['wifi_password'],password);self.assertEqual(len(read['vlan_details'][0]['networks']),2)
        changed=self.save('wifi',{'version':w['version'],'purpose':'保留密码'},w['id'],200)
        self.assertEqual(changed['wifi_password'],password)
        raw=self.client.get('/api/ledger/wifi/export').content;wb=load_workbook(BytesIO(raw))
        self.assertNotIn('名称',[c.value for c in wb['数据'][1]])
        self.assertFalse(any(c.value==password for cells in wb['数据'] for c in cells))
        imported='QA imported WiFi value'
        headers=[c.value for c in wb['数据'][1]]
        wb['数据'].cell(2,headers.index('密码')+1,imported)
        wb['数据'].cell(2,headers.index('认证方式')+1,'密码')
        out=BytesIO();wb.save(out)
        r=self.client.post('/api/ledger/wifi/imports',files={'file':('wifi.xlsx',out.getvalue())},data={'mode':'upsert'})
        self.assertEqual(r.status_code,201,r.text);batch=r.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
        self.assertNotIn(imported,str(staged['rows']))
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        changed=self.client.get('/api/ledger/item/'+w['id']).json();self.assertEqual(changed['wifi_password'],imported)
        changed=self.save('wifi',{'version':changed['version'],'auth_method':'none'},w['id'],200)
        self.assertIsNone(changed['wifi_password'])

    def test_numeric_sort_filters_and_device_credentials(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        for n in [100,2,10,*range(200,220)]:
            self.save('vlan',{'site_id':site,'vlan_number':n,'notes':'排序验证'})
        first=self.client.get('/api/ledger/vlan').json()
        self.assertEqual(first['page_size'],10)
        self.assertEqual(len(first['items']),10)
        self.assertEqual([v['vlan_number'] for v in first['items']][:3],[2,10,100])
        desc=self.client.get('/api/ledger/vlan?sort=vlan_desc&page=2&page_size=20').json()
        self.assertEqual([v['vlan_number'] for v in desc['items']],[100,10,2])
        self.assertEqual(len(self.client.get('/api/ledger/vlan?page_size=50').json()['items']),23)
        self.assertEqual(self.client.get('/api/ledger/vlan?page_size=30').status_code,422)
        self.assertEqual(self.client.get('/api/ledger/vlan?q=排序验证&site_id='+site).json()['total'],23)
        self.assertEqual(self.client.get('/api/ledger/vlan?site_id=missing').json()['total'],0)
        password='Device QA credential! '
        d=self.save('device',{'site_id':site,'installation_location':'主机房核心机柜','is_stacked':True,'login_username':'netadmin','login_password':password})
        self.assertTrue(d['is_stacked']);self.assertEqual(d['login_password'],password)
        fields=next(v['fields'] for v in self.client.get('/api/ledger/schema').json()['modules'] if v['kind']=='device')
        self.assertTrue({'management_vlan_id','owner_user_id','location_text','location_id'}.isdisjoint(f['name'] for f in fields))
        for path in ('/api/ledger/device','/api/ledger/options'):
            self.assertNotIn(password,self.client.get(path).text)
        with self.app.state.engine.begin() as conn:
            stored=conn.execute(sa.select(T['ledger_profiles'].c.data_json).where(T['ledger_profiles'].c.resource_id==d['id'])).scalar_one()
            self.assertNotIn(password,stored)
            conn.execute(T['resources'].update().where(T['resources'].c.id==d['id']).values(owner_user_id=self.it_a))
        d=self.save('device',{'version':d['version'],'notes':'更新测试'},d['id'],200)
        self.assertEqual(d['owner_user_id'],self.it_a);self.assertEqual(d['login_password'],password)
        wb=load_workbook(BytesIO(self.client.get('/api/ledger/device/export').content))
        headers=[c.value for c in wb['数据'][1]]
        self.assertIn('安装位置',headers);self.assertIn('密码',headers)
        self.assertIsNone(wb['数据'].cell(2,headers.index('密码')+1).value)
        imported='Device imported QA'
        wb['数据'].cell(2,headers.index('密码')+1,imported)
        out=BytesIO();wb.save(out)
        r=self.client.post('/api/ledger/device/imports',files={'file':('device.xlsx',out.getvalue())},data={'mode':'upsert'})
        self.assertEqual(r.status_code,201,r.text);batch=r.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertNotIn(imported,str(staged['rows']))
        self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.get('/api/ledger/item/'+d['id']).json()['login_password'],imported)

    def test_circuit_modes_encrypted_credentials_and_import(self):
        import csv
        from io import StringIO
        c=self.save('circuit',self.base('多IP线路',vendor_name='测试运营商',monthly_fee='899.50',support_phone='400-000-0000',ont_username='ont-admin',ont_password='ONT QA secret',lines=[
            {'line_type':'primary','number':'MAIN','downstream':'100'}, {'line_type':'gift','number':'GIFT'}],connections=[
            {'line_type':'primary','connection_method':'route','ip_address':'203.0.113.2','mask':'29','gateway':'203.0.113.1','interface':'GE1','purpose':'业务专线'},
            {'line_type':'primary','connection_method':'bridge','ip_address':'203.0.113.3','interface':'GE2','purpose':'桥接网络'},
            {'line_type':'gift','connection_method':'pppoe','dial_username':'dial-user','dial_password':'Dial QA secret ','ip_address':'invalid-hidden-ip','interface':'GE3','purpose':'普通上网'}]))
        self.assertEqual(len(c['connections']),3);self.assertEqual(c['connections'][2]['ip_address'],'')
        self.assertEqual(c['monthly_fee'],'899.50');self.assertNotIn('notes',c['lines'][0])
        self.assertEqual(c['connections'][2]['dial_password'],'Dial QA secret ')
        for path in ('/api/ledger/circuit','/api/ledger/options'):
            response=self.client.get(path).text
            self.assertNotIn('Dial QA secret',response);self.assertNotIn('ONT QA secret',response);self.assertNotIn('ciphertext',response)
        raw=self.client.get('/api/ledger/circuit/export').content;wb=load_workbook(BytesIO(raw))
        headers=[v.value for v in wb['数据'][1]];column=headers.index('线路接入明细·拨号密码')+1
        self.assertIn('保障电话',headers)
        self.assertIn('月租费用（元）',headers);self.assertNotIn('主线·月租费用（元）',headers);self.assertNotIn('主线·线路备注',headers)
        self.assertLess(headers.index('线路接入明细·拨号账号'),headers.index('线路接入明细·拨号密码'))
        self.assertLess(headers.index('线路接入明细·拨号密码'),headers.index('线路接入明细·用途'))
        self.assertEqual(wb['数据'].max_row,4)
        self.assertTrue(all(wb['数据'].cell(i,column).value is None for i in range(2,5)))
        wb['数据'].cell(4,column,'Imported dial QA');out=BytesIO();wb.save(out)
        r=self.client.post('/api/ledger/circuit/imports',files={'file':('circuit.xlsx',out.getvalue())},data={'mode':'upsert'})
        self.assertEqual(r.status_code,201,r.text);batch=r.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
        self.assertNotIn('Imported dial QA',str(staged))
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        c=self.client.get('/api/ledger/item/'+c['id']).json()
        self.assertEqual(c['connections'][2]['dial_password'],'Imported dial QA');self.assertEqual(c['ont_password'],'ONT QA secret')
        raw=self.client.get('/api/ledger/circuit/export').content
        batch=self.client.post('/api/ledger/circuit/imports',files={'file':('circuit.xlsx',raw)},data={'mode':'upsert'}).json()
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        c=self.client.get('/api/ledger/item/'+c['id']).json()
        self.assertEqual(c['connections'][2]['dial_password'],'Imported dial QA');self.assertEqual(len(c['connections']),3)
        c['connections'][2]['connection_method']='route'
        c=self.save('circuit',{'version':c['version'],'connections':c['connections']},c['id'],200)
        self.assertFalse(c['connections'][2]['dial_password'])
        self.save('circuit',{'version':c['version'],'connections':[{'line_type':'primary','connection_method':'invalid'}]},c['id'],422)
        self.save('circuit',{'version':c['version'],'connections':[{'line_type':'primary','connection_method':'route','ip_address':'bad'}]},c['id'],422)
        self.save('circuit',{'version':c['version'],'monthly_fee':'12.345'},c['id'],422)

    def test_aws_remote_vlan_notes_live_details_and_legacy(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        r=self.client.post('/api/data/sites',json={'name':'AWS区域'})
        self.assertEqual(r.status_code,201,r.text);remote=r.json()['id']
        local=self.save('vlan',{'site_id':site,'vlan_number':100,'purpose':'本地办公','networks':[{'cidr':'10.60.0.0/24'}]})
        peer=self.save('vlan',{'site_id':remote,'vlan_number':100,'purpose':'云端业务','networks':[{'cidr':'172.20.0.0/24'},{'cidr':'172.20.1.0/24'}]})
        vpn=self.save('vpn',{'site_id':site,'remote_site_id':remote,'local_vlan_ids':[local['id']],'remote_vlan_ids':[peer['id']],
            'remote_networks':[{'cidr':'172.30.0.0/16','notes':'手动补充 | 合作网段\n第二行'}]})
        self.assertEqual(len(vpn['remote_network_details']),3)
        self.assertEqual(vpn['remote_network_details'][0]['purpose'],'云端业务')
        self.assertEqual(vpn['local_network_details'][0]['purpose'],'本地办公')
        self.assertNotIn('notes',vpn['remote_network_details'][0])
        spec=next(v for v in self.client.get('/api/ledger/schema').json()['modules'] if v['kind']=='vpn')
        self.assertTrue({'purpose','local_network_notes','remote_network_notes'}.isdisjoint(f['name'] for f in spec['fields']))
        self.assertIn('172.30.0.0/16',self.client.get('/api/ledger/vpn').text)
        self.save('vpn',{'version':vpn['version'],'remote_vlan_ids':[local['id']]},vpn['id'],422)
        self.save('vlan',{'version':peer['version'],'networks':[{'cidr':'172.21.0.0/24'}]},peer['id'],409)
        self.save('vlan',{'version':peer['version'],'acknowledge_impacts':True,'networks':[{**peer['networks'][0],'cidr':'172.21.0.0/24'},peer['networks'][1]]},peer['id'],200)
        vpn=self.client.get('/api/ledger/item/'+vpn['id']).json()
        self.assertEqual(vpn['remote_network_details'][0]['cidr'],'172.21.0.0/24')
        raw=self.client.get('/api/ledger/vpn/export').content
        r=self.client.post('/api/ledger/vpn/imports',files={'file':('vpn.xlsx',raw)},data={'mode':'upsert'})
        self.assertEqual(r.status_code,201,r.text);batch=r.json()
        staged=self.client.get('/api/imports/batches/'+batch['id']).json()
        self.assertEqual(__import__('json').loads(staged['batch']['summary_json'])['errors'],0,staged)
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        roundtrip=self.client.get('/api/ledger/item/'+vpn['id']).json()
        self.assertEqual(roundtrip['remote_networks'][0]['notes'],'手动补充 | 合作网段\n第二行')
        with self.app.state.engine.begin() as conn:
            profile=conn.execute(sa.select(T['ledger_profiles'].c.data_json).where(T['ledger_profiles'].c.resource_id==vpn['id'])).scalar_one()
            old=__import__('json').loads(profile);old.pop('remote_networks');old['remote_cidrs']='192.0.2.0/24'
            conn.execute(T['ledger_profiles'].update().where(T['ledger_profiles'].c.resource_id==vpn['id']).values(data_json=__import__('json').dumps(old)))
        self.assertEqual(self.client.get('/api/ledger/item/'+vpn['id']).json()['remote_networks'][0]['cidr'],'192.0.2.0/24')

    def test_legacy_multi_ip_and_encrypted_dial_projection(self):
        import json
        from portal import ledger_secrets
        c=self.save('circuit',self.base('旧线路',lines=[{'line_type':'primary','number':'OLD'}]))
        with self.app.state.engine.begin() as conn:
            old={'vendor_name':'旧运营商','contact':'400-000-0000','lines':[
                {'line_type':'primary','number':'OLD','connection_method':'static','usable_ips':'203.0.113.2\n203.0.113.3','interface':'GE1','purpose':'旧专线'},
                {'line_type':'gift','connection_method':'pppoe','dial_username':'legacy-user','dial_password_ciphertext':ledger_secrets.encrypt(conn,'Legacy dial QA secret')} ]}
            conn.execute(T['ledger_profiles'].update().where(T['ledger_profiles'].c.resource_id==c['id']).values(data_json=json.dumps(old)))
        response=self.client.get('/api/ledger/circuit').text
        self.assertNotIn('Legacy dial QA secret',response);self.assertNotIn('ciphertext',response)
        obj=self.client.get('/api/ledger/item/'+c['id']).json()
        self.assertEqual(len(obj['connections']),3);self.assertEqual(obj['connections'][2]['dial_password'],'Legacy dial QA secret')
        self.assertEqual(obj['support_phone'],'400-000-0000');self.assertEqual(obj['connections'][0]['connection_method'],'route')
        changed=self.save('circuit',{'version':obj['version'],'enabled':False},obj['id'],200)
        self.assertFalse(changed['enabled']);self.assertEqual(changed['connections'][2]['dial_password'],'Legacy dial QA secret')

    def test_per_network_purpose_legacy_fallback_and_tunnel_states(self):
        import json
        v=self.save('vlan',self.base('分网段用途',vlan_number=380,networks=[
            {'cidr':'10.80.0.0/24','gateway':'10.80.0.1','purpose':'办公终端'},
            {'cidr':'10.81.0.0/24','gateway':'10.81.0.1','purpose':'生产设备'}]))
        vpn=self.save('vpn',self.base('双隧道',local_vlan_ids=[v['id']],tunnels=[
            {'remote_gateway':'192.0.2.1','registered_state':'up'},
            {'remote_gateway':'192.0.2.2','registered_state':'停用'}]))
        self.assertEqual([n['purpose'] for n in vpn['local_network_details']],['办公终端','生产设备'])
        self.assertEqual([t['registered_state'] for t in vpn['tunnels']],['up','disabled'])
        changed=[{**n,'purpose':'研发办公' if i==0 else n['purpose']} for i,n in enumerate(v['networks'])]
        v=self.save('vlan',{'version':v['version'],'networks':changed},v['id'],200)
        self.assertEqual(self.client.get('/api/ledger/item/'+vpn['id']).json()['local_network_details'][0]['purpose'],'研发办公')
        raw=self.client.get('/api/ledger/vlan/export').content
        batch=self.client.post('/api/ledger/vlan/imports',files={'file':('vlan.xlsx',raw)},data={'mode':'upsert'}).json()
        r=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(r.status_code,200,r.text)
        self.assertEqual([n['purpose'] for n in self.client.get('/api/ledger/item/'+v['id']).json()['networks']],['研发办公','生产设备'])
        self.save('vpn',self.base('非法状态',tunnels=[{'registered_state':'down'}]),status=422)
        with self.app.state.engine.begin() as conn:
            for key,update in [(v['id'],{'purpose':'旧版共同用途'}),(vpn['id'],None)]:
                stored=json.loads(conn.execute(sa.select(T['ledger_profiles'].c.data_json).where(T['ledger_profiles'].c.resource_id==key)).scalar_one())
                if update:
                    stored.update(update)
                    for n in stored['networks']:n.pop('purpose',None)
                else:stored['tunnels'][1]['registered_state']='unknown'
                conn.execute(T['ledger_profiles'].update().where(T['ledger_profiles'].c.resource_id==key).values(data_json=json.dumps(stored)))
        old=self.client.get('/api/ledger/item/'+v['id']).json()
        self.assertEqual([n['purpose'] for n in old['networks']],['旧版共同用途','旧版共同用途'])
        vpn=self.client.get('/api/ledger/item/'+vpn['id']).json()
        self.assertEqual(vpn['tunnels'][1]['registered_state'],'unknown')
        vpn=self.save('vpn',{'version':vpn['version'],'notes':'保留旧状态待确认'},vpn['id'],200)
        self.assertEqual(vpn['tunnels'][1]['registered_state'],'unknown')
        vpn=self.save('vpn',{'version':vpn['version'],'tunnels':[{'remote_gateway':'192.0.2.1','registered_state':'正常'},{'remote_gateway':'192.0.2.2','registered_state':'disabled'}]},vpn['id'],200)
        self.assertEqual(vpn['tunnels'][1]['registered_state'],'disabled')

    def test_independent_vlan_network_references_and_flat_rows(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        v=self.save('vlan',{'site_id':site,'vlan_number':390,'networks':[
            {'cidr':'10.90.0.0/24','gateway':'10.90.0.1','purpose':'办公网段'},
            {'cidr':'10.91.0.0/24','gateway':'10.91.0.1','purpose':'生产网段'}]})
        a,b=[v['id']+':'+n['network_id'] for n in v['networks']]
        listed=self.client.get('/api/ledger/vlan').json()
        self.assertEqual(listed['total'],2);self.assertEqual([x['id'] for x in listed['items']],[v['id'],v['id']])
        self.assertTrue(all(len(x['networks'])==1 for x in listed['items']))
        self.assertEqual(self.client.get('/api/ledger/vlan?q=办公网段').json()['total'],1)
        choices=self.client.get('/api/ledger/options').json()['vlan_network']
        self.assertEqual({x['id'] for x in choices},{a,b})
        wa=self.save('wifi',{'site_id':site,'ssid':'Only-Office','vlan_network_ids':[a]})
        wb=self.save('wifi',{'site_id':site,'ssid':'Only-Factory','vlan_network_ids':[b]})
        self.assertEqual(wa['vlan_ids'],[v['id']]);self.assertEqual(len(wa['vlan_details'][0]['networks']),1)
        self.assertEqual(wb['vlan_details'][0]['networks'][0]['purpose'],'生产网段')
        vpn=self.save('vpn',{'site_id':site,'local_vlan_network_ids':[a],'remote_site_id':site,'remote_vlan_network_ids':[b]})
        self.assertEqual([n['cidr'] for n in vpn['local_network_details']],['10.90.0.0/24'])
        self.assertEqual([n['cidr'] for n in vpn['remote_network_details']],['10.91.0.0/24'])
        self.save('wifi',{'site_id':site,'ssid':'Missing-Net','vlan_network_ids':[v['id']+':missing']},status=409)
        changed=self.save('vlan',{'version':v['version'],'acknowledge_impacts':True,'networks':[
            {**v['networks'][0],'cidr':'10.92.0.0/24','gateway':'10.92.0.1'},v['networks'][1],
            {'cidr':'10.93.0.0/24','purpose':'未选择的新网段'}]},v['id'],200)
        self.assertEqual(changed['networks'][0]['network_id'],v['networks'][0]['network_id'])
        wa=self.client.get('/api/ledger/item/'+wa['id']).json()
        self.assertEqual([n['cidr'] for n in wa['vlan_details'][0]['networks']],['10.92.0.0/24'])
        self.save('vlan',{'version':changed['version'],'acknowledge_impacts':True,'networks':changed['networks'][1:]},v['id'],409)
        raw=self.client.get('/api/ledger/wifi/export').content
        wbfile=load_workbook(BytesIO(raw));cells=str(list(wbfile['数据'].values))
        self.assertIn('10.92.0.0/24',cells)
        self.assertNotIn('default@390@10.92.0.0/24',cells)
        batch=self.client.post('/api/ledger/wifi/imports',files={'file':('wifi.xlsx',raw)},data={'mode':'upsert'}).json()
        result=self.client.post('/api/imports/batches/'+batch['id']+'/commit',json={});self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(self.client.get('/api/ledger/item/'+wa['id']).json()['vlan_network_ids'],[a])
        wb=self.client.get('/api/ledger/item/'+wb['id']).json()
        self.save('wifi',{'version':wb['version'],'vlan_network_ids':[]},wb['id'],200)
        self.save('vpn',{'version':vpn['version'],'remote_vlan_network_ids':[]},vpn['id'],200)
        self.save('vlan',{'version':changed['version'],'acknowledge_impacts':True,'networks':[changed['networks'][0],changed['networks'][2]]},v['id'],200)

    def test_delete_only_one_network_retains_vlan_and_sibling_references(self):
        v=self.save('vlan',self.base('删除网段测试',vlan_number=395,notes='保留VLAN备注',networks=[
            {'cidr':'10.95.0.0/24','purpose':'待删除'}, {'cidr':'10.96.0.0/24','purpose':'需要保留'}]))
        first,second=v['networks'];path='/api/ledger/vlan/'+v['id']+'/networks/'
        w=self.save('wifi',self.base('保留引用',ssid='Keep-Sibling',vlan_network_ids=[v['id']+':'+second['network_id']]))
        r=self.client.request('DELETE','/api/ledger/item/'+v['id'],json={'version':v['version'],'confirm':v['id'],'reason':'禁止整VLAN删除'})
        self.assertEqual(r.status_code,409,r.text)
        r=self.client.request('DELETE',path+first['network_id'],json={'version':v['version'],'confirm':first['network_id'],'reason':'仅删除当前网段'})
        self.assertEqual(r.status_code,200,r.text);result=r.json()
        self.assertEqual(result['remaining_networks'],1)
        kept=self.client.get('/api/ledger/item/'+v['id']).json()
        self.assertEqual(kept['vlan_number'],395);self.assertEqual(kept['notes'],'保留VLAN备注')
        self.assertEqual(kept['networks'],[second])
        self.assertEqual(self.client.get('/api/ledger/item/'+w['id']).json()['vlan_network_ids'],[v['id']+':'+second['network_id']])
        r=self.client.request('DELETE',path+second['network_id'],json={'version':v['version'],'confirm':second['network_id'],'reason':'旧版本'})
        self.assertEqual(r.status_code,409,r.text)
        r=self.client.request('DELETE',path+second['network_id'],json={'version':kept['version'],'confirm':second['network_id'],'reason':'不能删除被引用网段'})
        self.assertEqual(r.status_code,409,r.text)
        self.save('wifi',{'version':w['version'],'vlan_network_ids':[]},w['id'],200)
        r=self.client.request('DELETE',path+second['network_id'],json={'version':kept['version'],'confirm':second['network_id'],'reason':'最后一个网段也保留父记录'})
        self.assertEqual(r.status_code,200,r.text)
        empty=self.client.get('/api/ledger/item/'+v['id']).json()
        self.assertEqual(empty['networks'],[]);self.assertEqual(empty['vlan_number'],395)
        with self.app.state.engine.connect() as conn:
            self.assertIsNotNone(conn.execute(sa.select(T['vlans']).where(T['vlans'].c.resource_id==v['id'])).first())
            self.assertIsNotNone(conn.execute(sa.select(T['ledger_profiles']).where(T['ledger_profiles'].c.resource_id==v['id'])).first())
        self.login('employee-a')
        r=self.client.request('DELETE',path+second['network_id'],json={'version':empty['version'],'confirm':second['network_id'],'reason':'越权'})
        self.assertEqual(r.status_code,403,r.text)

    def test_whole_vlan_delete_blocked_and_delete_permission_does_not_require_edit(self):
        v=self.save('vlan',self.base('专用删除权限',vlan_number=396,networks=[{'cidr':'10.97.0.0/24'}]))
        r=self.client.request('DELETE','/api/ledger/item/'+v['id'],json={'version':v['version'],'confirm':v['id'],'reason':'阻止旧接口误删'})
        self.assertEqual(r.status_code,409,r.text)
        role=self.client.post('/api/roles',json={'code':'network_delete_only','name':'仅网段删除','permissions':[
            {'code':'resource.view','scope':'all'},{'code':'resource.delete','scope':'all'}]})
        self.assertEqual(role.status_code,201,role.text)
        user=next(u for u in self.client.get('/api/users').json()['items'] if u['id']==self.it_a)
        r=self.client.put('/api/users/'+self.it_a,json={'version':user['version'],'role_ids':[role.json()['id']]})
        self.assertEqual(r.status_code,200,r.text);self.login('it-a')
        n=v['networks'][0]
        r=self.client.request('DELETE','/api/ledger/vlan/'+v['id']+'/networks/'+n['network_id'],json={'version':v['version'],'confirm':n['network_id'],'reason':'只有删除权限也可删除网段'})
        self.assertEqual(r.status_code,200,r.text)
        empty=self.client.get('/api/ledger/item/'+v['id']).json();self.assertEqual(empty['networks'],[])
        r=self.client.request('DELETE','/api/ledger/item/'+v['id'],json={'version':empty['version'],'confirm':v['id'],'reason':'明确删除空VLAN'})
        self.assertEqual(r.status_code,200,r.text)

    def test_selected_static_network_ipam_and_computer_link(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        vlan=self.save('vlan',{'site_id':site,'vlan_number':398,'networks':[
            {'cidr':'10.98.0.0/29','gateway':'10.98.0.1','purpose':'静态终端','ipam_enabled':True},
            {'cidr':'10.98.1.0/29','purpose':'DHCP终端','ipam_enabled':False}]})
        managed=vlan['id']+':'+vlan['networks'][0]['network_id']
        options=self.client.get('/api/ipam/options').json()
        self.assertEqual([v['id'] for v in options['networks']],[managed])
        computer=self.client.post('/api/computers',json={'device_type':'台式机','site_id':site,
            'department_id':self.department,'user_id':self.employee,'location':'研发办公室','adapters':[{
                'adapter_type':'wired','mac_address':'001122334455'}]} )
        self.assertEqual(computer.status_code,201,computer.text);computer=computer.json();adapter=computer['adapters'][0]
        self.assertEqual(computer['device_name'],computer['asset_code']);self.assertEqual(adapter['adapter_name'],'有线网卡')
        listing=self.client.get('/api/ipam/addresses',params={'network_key':managed}).json()
        self.assertEqual(listing['stats']['total'],6);self.assertEqual(listing['stats']['reserved'],1)
        self.assertEqual(listing['page_size'],10)
        saved=self.client.post('/api/ipam/records',json={'network_key':managed,'ip_address':'10.98.0.2','occupancy_status':'occupied',
            'verification_status':'confirmed','computer_id':computer['id'],'adapter_id':adapter['id'],'device_type':'打印机','notes':'研发固定地址'})
        self.assertEqual(saved.status_code,200,saved.text);saved=saved.json()
        self.assertEqual(saved['mac_address'],'00:11:22:33:44:55');self.assertEqual(saved['department_id'],self.department);self.assertEqual(saved['device_type'],'台式机')
        filtered=self.client.get('/api/ipam/addresses',params={'network_key':managed,'q':'00:11:22:33:44:55'}).json()
        self.assertEqual(filtered['total'],1);self.assertEqual(filtered['items'][0]['computer_name'],computer['asset_code'])
        linked_computer=self.client.get('/api/computers/'+computer['id']).json()
        self.assertEqual([v['ip_address'] for v in linked_computer['ip_addresses']],['10.98.0.2'])
        deleted=self.client.request('DELETE','/api/computers/'+computer['id'],json={'version':computer['version'],'reason':'验证引用保护'})
        self.assertEqual(deleted.status_code,409,deleted.text)
        blocked=self.save('vlan',{'version':vlan['version'],'networks':[{**vlan['networks'][0],'ipam_enabled':False},vlan['networks'][1]]},vlan['id'],409)
        released=self.client.post('/api/ipam/records',json={'network_key':managed,'ip_address':'10.98.0.2','version':saved['version'],
            'occupancy_status':'available','verification_status':'pending','notes':'已释放'})
        self.assertEqual(released.status_code,200,released.text)
        history=self.client.get('/api/ipam/records/'+saved['id']+'/history').json()['items'];self.assertEqual(len(history),2)
        deleted=self.client.request('DELETE','/api/computers/'+computer['id'],json={'version':computer['version'],'reason':'测试删除'})
        self.assertEqual(deleted.status_code,200,deleted.text)
        self.assertEqual(self.client.get('/api/computers/'+computer['id']).status_code,404)

    def test_computer_entry_claims_ip_and_tracks_accessory_changes(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        vlan=self.save('vlan',{'site_id':site,'vlan_number':399,'networks':[{'cidr':'10.99.0.0/29','gateway':'10.99.0.1','purpose':'电脑','ipam_enabled':True}]})
        accepted=f'{date.today().year-6}-{date.today().month:02d}'
        computer=self.client.post('/api/computers',json={'device_type':'笔记本','site_id':site,'user_id':self.employee,
            'acceptance_month':accepted,
            'ip_assignments':[{'ip_address':'10.99.0.2','adapter_mac':'001122334466'}],'adapters':[{'adapter_type':'wired','mac_address':'001122334466'}],
            'accessories':[{'accessory_type':'显示器','brand_model':'Dell P2422H','serial_number':'MON-01'}]})
        self.assertEqual(computer.status_code,201,computer.text);computer=computer.json()
        self.assertEqual(computer['acceptance_month'],accepted+'-01');self.assertEqual(computer['usage_months'],72);self.assertEqual(computer['usage_text'],'6年0个月');self.assertTrue(computer['is_over_five_years'])
        self.assertEqual(computer['department_id'],self.department)
        self.assertEqual([v['ip_address'] for v in computer['ip_addresses']],['10.99.0.2'])
        self.assertEqual(computer['accessories'][0]['accessory_type'],'显示器')
        network=vlan['id']+':'+vlan['networks'][0]['network_id']
        address=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.99.0.2'}).json()['items'][0]
        self.assertEqual(address['computer_id'],computer['id']);self.assertEqual(address['occupancy_status'],'occupied')
        updated=self.client.put('/api/computers/'+computer['id'],json={'version':computer['version'],'device_type':'笔记本','site_id':site,
            'user_id':self.other,'ip_assignments':[{'ip_address':'10.99.0.2','adapter_mac':'001122334466'}],'adapters':[{'id':computer['adapters'][0]['id'],'adapter_type':'wired','mac_address':'001122334466'}],
            'accessories':[{'accessory_type':'键盘','brand_model':'Logitech'}]})
        self.assertEqual(updated.status_code,200,updated.text)
        self.assertEqual(updated.json()['acceptance_month'],accepted+'-01')
        history=self.client.get('/api/computers/'+computer['id']+'/history').json()['items']
        self.assertIn('使用人：',history[0]['change_summary']);self.assertIn('配件：',history[0]['change_summary']);self.assertIn('→',history[0]['change_summary'])

    def test_shared_computer_allows_no_person_or_department(self):
        site=self.client.get('/api/lookups').json()['sites'][0]['id']
        vlan=self.save('vlan',{'site_id':site,'vlan_number':402,'networks':[{'cidr':'10.102.0.0/29','purpose':'公共终端','ipam_enabled':True}]})
        response=self.client.post('/api/computers',json={'device_type':'台式机','assignment_type':'shared','site_id':site,'location':'公共会议室',
            'ip_assignments':[{'ip_address':'10.102.0.2','adapter_mac':'001122334477'}],'adapters':[{'adapter_type':'wired','mac_address':'001122334477'}]})
        self.assertEqual(response.status_code,201,response.text);computer=response.json()
        self.assertEqual(computer['assignment_type'],'shared');self.assertIsNone(computer['user_id']);self.assertIsNone(computer['department_id'])
        network=vlan['id']+':'+vlan['networks'][0]['network_id']
        record=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.102.0.2'}).json()['items'][0]
        self.assertEqual(record['occupancy_status'],'occupied');self.assertEqual(record['computer_id'],computer['id']);self.assertIsNone(record['user_id'])
        public=self.client.post('/api/ipam/records',json={'network_key':network,'ip_address':'10.102.0.3','occupancy_status':'occupied',
            'verification_status':'pending','assignment_type':'shared','device_type':'打印机'})
        self.assertEqual(public.status_code,200,public.text);self.assertIsNone(public.json()['user_id']);self.assertIsNone(public.json()['department_id']);self.assertIsNone(public.json()['location'])

    def test_ipam_and_computer_column_templates_import_export_and_unlinked_filter(self):
        site=self.client.get('/api/lookups').json()['sites'][0];department=self.client.get('/api/lookups').json()['departments'][0]
        vlan=self.save('vlan',{'site_id':site['id'],'vlan_number':401,'networks':[{'cidr':'10.101.0.0/29','gateway':'10.101.0.1','purpose':'导入测试','ipam_enabled':True}]})
        network=vlan['id']+':'+vlan['networks'][0]['network_id']
        response=self.client.get('/api/computer-transfers/template');self.assertEqual(response.status_code,200,response.text)
        wb=load_workbook(BytesIO(response.content));self.assertEqual([v.value for v in wb['数据'][1]],['资产编号','设备类型','验收日期','归属类型','区域','部门','使用人','位置','操作系统','状态','网卡类型','MAC','关联 IP','配件类型','配件备注','电脑备注'])
        accepted=f'{date.today().year-6}-{date.today().month:02d}'
        wb['数据'].append(['PC-IMPORT-01','笔记本',accepted,'个人使用',site['name'],department['name'],'employee-a','研发区','Windows 11','在用','有线','00:11:22:AA:BB:CC','10.101.0.2','显示器','24寸','测试电脑'])
        content=BytesIO();wb.save(content);response=self.client.post('/api/computer-transfers/import',files={'file':('computers.xlsx',content.getvalue())})
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()['count'],1)
        computer=next(v for v in self.client.get('/api/computers').json()['items'] if v['asset_code']=='PC-IMPORT-01');self.assertEqual(computer['ip_addresses'][0]['ip_address'],'10.101.0.2')
        self.assertEqual(computer['acceptance_month'],accepted+'-01');self.assertEqual(computer['usage_months'],72);self.assertTrue(computer['is_over_five_years'])
        exported_computers=self.client.get('/api/computer-transfers/export');self.assertEqual(exported_computers.status_code,200)
        response=self.client.post('/api/computer-transfers/import',files={'file':('computer-export.xlsx',exported_computers.content)})
        self.assertEqual(response.status_code,200,response.text)

        response=self.client.get('/api/ipam-transfers/template');self.assertEqual(response.status_code,200,response.text)
        wb=load_workbook(BytesIO(response.content));self.assertEqual([v.value for v in wb['数据'][1]],['区域','VLAN ID','网段','IP 地址','归属类型','部门','使用人','占用设备','MAC地址','位置','备注','状态','确认状态'])
        wb['数据'].append([site['name'],401,'10.101.0.0/29','10.101.0.3','个人使用',department['name'],'employee-a','打印机','00:11:22:AA:BB:DD','打印区','导入记录','已占用','待确认'])
        wb['数据'].append([site['name'],401,'10.101.0.0/29','10.101.0.4','个人使用',department['name'],'employee-b','台式机','00:11:22:AA:BB:EE','办公区','待关联','已占用','待确认'])
        wb['数据'].append([site['name'],401,'10.101.0.0/29','10.101.0.5','','','','','','','','',''])
        content=BytesIO();wb.save(content);response=self.client.post('/api/ipam-transfers/import',files={'file':('ip.xlsx',content.getvalue())})
        self.assertEqual(response.status_code,200,response.text)
        listing=self.client.get('/api/ipam/addresses',params={'network_key':network,'asset_link':'unlinked'}).json()
        self.assertEqual(listing['total'],1);self.assertEqual(listing['items'][0]['ip_address'],'10.101.0.4');self.assertEqual(listing['items'][0]['consistency_status'],'pending')
        non_computer=self.client.get('/api/ipam/addresses',params={'network_key':network,'asset_link':'non_computer'}).json();self.assertEqual(non_computer['items'][0]['ip_address'],'10.101.0.3')
        available=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.5'}).json()['items'][0];self.assertEqual(available['occupancy_status'],'available');self.assertEqual(available['verification_status'],'pending')
        linked=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.2'}).json()['items'][0]
        changed=self.client.post('/api/ipam/records',json={'network_key':network,'ip_address':'10.101.0.2','version':linked['version'],'occupancy_status':'occupied','verification_status':'confirmed','computer_id':computer['id'],'adapter_id':computer['adapters'][0]['id'],'location':'错误位置'})
        self.assertEqual(changed.status_code,200,changed.text);conflict=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.2'}).json()['items'][0];self.assertEqual(conflict['consistency_status'],'conflict');self.assertEqual(conflict['location'],'错误位置')
        reconciled=self.client.post('/api/ipam/records/'+conflict['id']+'/reconcile',json={'version':conflict['version'],'action':'from_computer'});self.assertEqual(reconciled.status_code,200,reconciled.text)
        self.assertEqual(self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.2'}).json()['items'][0]['consistency_status'],'consistent')
        mismatch=listing['items'][0]
        rejected=self.client.post('/api/ipam/records',json={'network_key':network,'ip_address':mismatch['ip_address'],'version':mismatch['version'],'occupancy_status':'occupied','verification_status':'pending','computer_id':computer['id'],'adapter_id':computer['adapters'][0]['id']})
        self.assertEqual(rejected.status_code,409,rejected.text)
        created=self.client.post('/api/ipam/records',json={'network_key':network,'ip_address':'10.101.0.6','occupancy_status':'occupied','verification_status':'pending','user_id':self.other,'device_type':'台式机','mac_address':'00:11:22:AA:BB:CC','location':'原位置'});self.assertEqual(created.status_code,200,created.text)
        candidate=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.6'}).json()['items'][0]
        self.assertEqual(candidate['asset_link_status'],'candidate');self.assertEqual(candidate['candidate_asset_code'],'PC-IMPORT-01')
        self.assertEqual(candidate['location'],'原位置');self.assertEqual(candidate['consistency_status'],'pending')
        reconciled=self.client.post('/api/ipam/records/'+candidate['id']+'/reconcile',json={'version':candidate['version'],'action':'to_computer'});self.assertEqual(reconciled.status_code,200,reconciled.text)
        reconciled_computer=self.client.get('/api/computers/'+computer['id']).json()
        self.assertEqual(reconciled_computer['assignment_type'],'personal');self.assertEqual(reconciled_computer['user_id'],self.other)
        self.assertEqual(reconciled_computer['display_user'],'employee-b');self.assertEqual(reconciled_computer['device_type'],'台式机')
        options=self.client.get('/api/ipam/options').json();option=next(v for v in options['computers'] if v['id']==computer['id'])
        self.assertEqual(option['display_user'],'employee-b')
        linked_record=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.101.0.6'}).json()['items'][0]
        self.assertEqual(linked_record['asset_link_status'],'linked');self.assertEqual(linked_record['consistency_status'],'consistent')
        exported=self.client.get('/api/ipam-transfers/export',params={'network_key':network});self.assertEqual(exported.status_code,200,exported.text)
        exported_rows=list(load_workbook(BytesIO(exported.content),read_only=True)['数据'].iter_rows(min_row=2,values_only=True))
        self.assertEqual([value[3] for value in exported_rows],['10.101.0.2','10.101.0.3','10.101.0.4','10.101.0.5','10.101.0.6'])
        exported_available=next(value for value in exported_rows if value[3]=='10.101.0.5')
        self.assertEqual(exported_available[11:13],('空闲','待确认'));self.assertIsNone(exported_available[4])

    def test_ipam_import_allows_missing_and_departed_people_for_review(self):
        lookups=self.client.get('/api/lookups').json();site=lookups['sites'][0];department=lookups['departments'][0]
        vlan=self.save('vlan',{'site_id':site['id'],'vlan_number':403,'networks':[{'cidr':'10.103.0.0/29','gateway':'10.103.0.1','purpose':'人员核对','ipam_enabled':True}]})
        network=vlan['id']+':'+vlan['networks'][0]['network_id']
        with self.app.state.engine.begin() as conn:
            conn.execute(T['users'].update().where(T['users'].c.id==self.other).values(account_state='disabled',updated_at=business.utcnow()))

        template=self.client.get('/api/ipam-transfers/template');self.assertEqual(template.status_code,200,template.text)
        wb=load_workbook(BytesIO(template.content));sheet=wb['数据']
        sheet.append([site['name'],403,'10.103.0.0/29','10.103.0.2','个人使用',department['name'],'系统外人员','打印机','00:11:22:AA:CC:01','前台','','已占用','已确认'])
        sheet.append([site['name'],403,'10.103.0.0/29','10.103.0.3','个人使用',department['name'],'employee-b','台式机','00:11:22:AA:CC:02','旧工位','','已占用','已确认'])
        content=BytesIO();wb.save(content)
        imported=self.client.post('/api/ipam-transfers/import',files={'file':('people-review.xlsx',content.getvalue())})
        self.assertEqual(imported.status_code,200,imported.text);self.assertEqual(imported.json()['count'],2)
        self.assertEqual(imported.json()['people_review'],{'departed':1,'missing':1})

        missing=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.103.0.2'}).json()['items'][0]
        self.assertEqual(missing['display_user'],'系统外人员');self.assertEqual(missing['person_status'],'missing')
        self.assertEqual(missing['verification_status'],'pending');self.assertEqual(missing['consistency_status'],'pending')
        self.assertIn('人员不存在',missing['consistency_issues'])
        departed=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.103.0.3'}).json()['items'][0]
        self.assertEqual(departed['display_user'],'employee-b');self.assertEqual(departed['person_status'],'departed')
        self.assertEqual(departed['verification_status'],'pending');self.assertEqual(departed['consistency_status'],'pending')
        self.assertIn('人员已离职',departed['consistency_issues'])

        exported=self.client.get('/api/ipam-transfers/export',params={'network_key':network});self.assertEqual(exported.status_code,200,exported.text)
        exported_rows=list(load_workbook(BytesIO(exported.content),read_only=True)['数据'].iter_rows(min_row=2,values_only=True))
        self.assertIn('系统外人员',[value[6] for value in exported_rows]);self.assertIn('employee-b',[value[6] for value in exported_rows])

    def test_shared_ip_keeps_independent_management_department(self):
        lookups=self.client.get('/api/lookups').json();site=lookups['sites'][0];department=lookups['departments'][0]
        vlan=self.save('vlan',{'site_id':site['id'],'vlan_number':404,'networks':[{'cidr':'10.104.0.0/29','gateway':'10.104.0.1','purpose':'公共设备部门','ipam_enabled':True}]})
        network=vlan['id']+':'+vlan['networks'][0]['network_id']
        template=self.client.get('/api/ipam-transfers/template');wb=load_workbook(BytesIO(template.content));sheet=wb['数据']
        sheet.append([site['name'],404,'10.104.0.0/29','10.104.0.2','公共设备',department['name'],'','打印机','00:11:22:AA:DD:01','文印室','','已占用','已确认'])
        sheet.append([site['name'],404,'10.104.0.0/29','10.104.0.3','公共设备','待核对部门','','门禁','00:11:22:AA:DD:02','大门','','已占用','已确认'])
        content=BytesIO();wb.save(content)
        imported=self.client.post('/api/ipam-transfers/import',files={'file':('shared-departments.xlsx',content.getvalue())})
        self.assertEqual(imported.status_code,200,imported.text);self.assertEqual(imported.json()['department_review'],{'missing':1})

        known=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.104.0.2'}).json()['items'][0]
        self.assertEqual(known['department_id'],department['id']);self.assertEqual(known['department_name'],department['name'])
        self.assertEqual(known['department_status'],'');self.assertEqual(known['consistency_status'],'consistent')
        missing=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.104.0.3'}).json()['items'][0]
        self.assertIsNone(missing['department_id']);self.assertEqual(missing['department_name_snapshot'],'待核对部门')
        self.assertEqual(missing['department_name'],'待核对部门');self.assertEqual(missing['department_status'],'missing')
        self.assertEqual(missing['verification_status'],'pending');self.assertEqual(missing['consistency_status'],'pending')
        self.assertIn('部门不存在',missing['consistency_issues'])

        exported=self.client.get('/api/ipam-transfers/export',params={'network_key':network});self.assertEqual(exported.status_code,200,exported.text)
        rows=list(load_workbook(BytesIO(exported.content),read_only=True)['数据'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(next(value for value in rows if value[3]=='10.104.0.3')[5],'待核对部门')
        corrected=self.client.post('/api/ipam/records',json={'network_key':network,'ip_address':'10.104.0.3','version':missing['version'],
            'occupancy_status':'occupied','verification_status':'confirmed','assignment_type':'shared','department_id':department['id'],
            'department_name_snapshot':'','device_type':'门禁','mac_address':'00:11:22:AA:DD:02','location':'大门'})
        self.assertEqual(corrected.status_code,200,corrected.text);self.assertEqual(corrected.json()['department_id'],department['id'])
        corrected_view=self.client.get('/api/ipam/addresses',params={'network_key':network,'q':'10.104.0.3'}).json()['items'][0]
        self.assertEqual(corrected_view['department_name'],department['name']);self.assertEqual(corrected_view['department_status'],'')
        self.assertEqual(corrected_view['verification_status'],'confirmed')
